"""Year-end compensation: employees per calendar year, spreadsheet import, the 1604-C alphalist workbook, 2316s,
and a firm-wide 1604-C checklist."""

import asyncio
import io
import json
import uuid
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_current_user, get_db
from app.models.accounting import AACompany, AAEmployee, AAOutput, AAPeriod
from app.models.user import User
from app.routers.accounting_scope import MANAGER_ROLES, company_or_404, scope_for
from app.services.accounting import audit
from app.services.accounting.compute.alphalist import (
    FIELD_LABELS, MonthRemittance, build_alphalist, employee_from_json,
)
from app.services.accounting.compute.due_dates import month_start
from app.services.accounting.compute.money import ZERO, c2
from app.services.spreadsheet_safe import csv_safe as _safe  # imported sheet text is untrusted: never a formula
from app.services.accounting.filing import credited_output_id
from app.services.accounting.payroll_import import IMPORT_FIELDS, apply_mapping, parse_sheet, suggest_mapping
from app.schemas.accounting import CertificateOut, CertificatePreviewIn, CertificatesIn, PdfFieldOut
from app.services.accounting.certificate_edits import apply_2316, fields_2316
from app.services.accounting.pdf_2316 import build_2316_zip, fill_2316

router = APIRouter()
MAX_SHEET_BYTES = 10 * 1024 * 1024


class EmployeeOut(BaseModel):
    id: uuid.UUID
    data: dict
    schedule: int
    gross: Decimal
    taxable: Decimal
    tax_due: Decimal
    withheld: Decimal              # present employer, Jan–Nov
    december: Decimal              # 17a: still to withhold
    refund: Decimal                # 17b: over-withheld
    substituted: bool


class EmployeesOut(BaseModel):
    year: int
    employees: list[EmployeeOut]
    warnings: list[str]
    checks: list[dict]
    remittances: list[dict]
    fields: dict[str, str]
    due: date


class PreviewOut(BaseModel):
    headers: list[str]
    sample: list[list[str]]
    rows: int
    mapping: dict[str, str]
    fields: dict[str, str]


class ImportOut(EmployeesOut):
    imported: int
    errors: list[str]


class StatusRow(BaseModel):
    company_id: uuid.UUID
    company: str
    employees: int
    issues: int
    status: str                    # no_payroll | needs_review | ready


def _require_manager(user: User) -> None:
    if user.role not in MANAGER_ROLES:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Only owners and managers can change employee records")


def _clean(data: dict) -> dict:
    """Keep known fields only, as JSON-safe values."""
    out = {}
    for key, value in data.items():
        if key in FIELD_LABELS and value not in (None, ""):
            out[key] = value if isinstance(value, bool) else str(value)
    return out


async def remittances_for(db: AsyncSession, company: AACompany, year: int) -> list[MonthRemittance]:
    """Each calendar month's 1601-C from the quarter folders: the filed version, else the latest."""
    fye = company.fiscal_year_end_month
    periods = (await db.execute(select(AAPeriod).where(
        AAPeriod.company_id == company.id, AAPeriod.quarter >= 1, AAPeriod.year.in_((year, year + 1)),
    ))).scalars().all()
    out: dict[int, MonthRemittance] = {}
    for p in periods:
        latest = (await db.execute(select(AAOutput).where(AAOutput.period_id == p.id)
                                   .order_by(AAOutput.version.desc()).limit(1))).scalars().first()
        for m in (1, 2, 3):
            first = month_start(p.year, p.quarter, fye, m)
            if first.year != year:
                continue
            form = f"1601C-M{m}"
            entry = (p.filed_forms or {}).get(form)
            filed_id = credited_output_id(p.filed_forms or {}, form, None)
            chosen = (await db.get(AAOutput, uuid.UUID(filed_id))) if filed_id else latest
            values = (chosen.snapshot.get("forms", {}).get(form) if chosen else None) or {}
            if not values:
                continue

            def num(item):
                raw = values.get(item)
                return c2(Decimal(str(raw))) if raw not in (None, "") else ZERO

            out[first.month] = MonthRemittance(
                first.month, (entry or {}).get("paid_on") or (entry or {}).get("filed_at", "")[:10] if entry else "",
                num("14"), num("25"), num("26"), num("35"), num("36"), "filed" if entry else "latest")
    return list(out.values())


async def _rows(db: AsyncSession, company_id: uuid.UUID, year: int) -> list[AAEmployee]:
    return list((await db.execute(select(AAEmployee).where(
        AAEmployee.company_id == company_id, AAEmployee.year == year).order_by(AAEmployee.created_at)
    )).scalars().all())


async def _build(db: AsyncSession, company: AACompany, year: int):
    rows = await _rows(db, company.id, year)
    employees = [employee_from_json(str(r.id), r.data) for r in rows]
    remittances = await remittances_for(db, company, year)
    tabs, computed, warnings = build_alphalist(employees, year, remittances)
    return rows, tabs, computed, warnings


async def _out(db: AsyncSession, company: AACompany, year: int) -> EmployeesOut:
    rows, tabs, computed, warnings = await _build(db, company, year)
    by_id = {c.employee.id: c for c in computed}
    employees = []
    for r in rows:
        c = by_id[str(r.id)]
        employees.append(EmployeeOut(
            id=r.id, data=r.data, schedule=c.schedule, gross=c.form["19"], taxable=c.form["23"],
            tax_due=c.form["24"], withheld=c.alpha["15b"], december=c.alpha["17a"], refund=c.alpha["17b"],
            substituted=c.form["substituted"]))
    employees.sort(key=lambda e: (str(e.data.get("last_name", "")).upper(), str(e.data.get("first_name", "")).upper()))
    return EmployeesOut(year=year, employees=employees, warnings=warnings, checks=tabs["Checks"],
                        remittances=tabs["Remittances"], fields=FIELD_LABELS, due=date(year + 1, 1, 31))


@router.get("/companies/{company_id}/employees/{year}", response_model=EmployeesOut)
async def list_employees(company_id: uuid.UUID, year: int, user: User = Depends(get_current_user),
                         db: AsyncSession = Depends(get_db)):
    company = await company_or_404(db, scope_for(user), company_id)
    return await _out(db, company, year)


async def _read_sheet(file: UploadFile) -> tuple[list[str], list[list[str]]]:
    data = await file.read(MAX_SHEET_BYTES + 1)
    if len(data) > MAX_SHEET_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "The sheet is larger than 10 MB")
    try:
        return await asyncio.get_running_loop().run_in_executor(None, parse_sheet, data, file.filename or "")
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from None
    except Exception:  # a corrupt workbook
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Couldn't read that file as Excel or CSV") from None


@router.post("/companies/{company_id}/employees/{year}/import/preview", response_model=PreviewOut)
async def preview_import(company_id: uuid.UUID, year: int, file: UploadFile, user: User = Depends(get_current_user),
                         db: AsyncSession = Depends(get_db)):
    company = await company_or_404(db, scope_for(user), company_id)
    headers, rows = await _read_sheet(file)
    return PreviewOut(headers=headers, sample=rows[:5], rows=len(rows),
                      mapping=suggest_mapping(headers, company.payroll_mapping or {}), fields=IMPORT_FIELDS)


@router.post("/companies/{company_id}/employees/{year}/import", response_model=ImportOut)
async def import_employees(company_id: uuid.UUID, year: int, file: UploadFile, mapping: str = Form(...),
                           user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    """Replace the year's employees with the sheet's rows, and remember the column mapping for this client."""
    _require_manager(user)
    ent = scope_for(user)
    company = await company_or_404(db, ent, company_id)
    try:
        chosen = {str(k): str(v) for k, v in json.loads(mapping).items() if v}
    except (ValueError, AttributeError):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid column mapping") from None
    if not ({"full_name", "last_name"} & set(chosen.values())):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Map a name column (full name or last name).")
    headers, rows = await _read_sheet(file)
    records, errors = apply_mapping(headers, rows, chosen)
    if not records:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No employee rows found with that mapping."
                            + (f" {errors[0]}" if errors else ""))
    replaced = (await db.execute(select(func.count(AAEmployee.id)).where(
        AAEmployee.company_id == company.id, AAEmployee.year == year))).scalar_one()
    await db.execute(delete(AAEmployee).where(AAEmployee.company_id == company.id, AAEmployee.year == year))
    audit.record(db, label=f"{year} employees",
                 ent=ent, company_id=company.id, user=user, entity="employee", entity_id=f"{year}",
                 action="import", before={"employees": replaced},
                 after={"employees": len(records), "file": (file.filename or "")[:120]}, always=True)
    for data in records:
        db.add(AAEmployee(enterprise_id=ent, company_id=company.id, year=year, data=_clean(data)))
    company.payroll_mapping = {**(company.payroll_mapping or {}), **chosen}
    await db.commit()
    out = await _out(db, company, year)
    return ImportOut(**out.model_dump(), imported=len(records), errors=errors)


@router.post("/companies/{company_id}/employees/{year}", response_model=EmployeesOut,
             status_code=status.HTTP_201_CREATED)
async def add_employee(company_id: uuid.UUID, year: int, body: dict, user: User = Depends(get_current_user),
                       db: AsyncSession = Depends(get_db)):
    _require_manager(user)
    ent = scope_for(user)
    company = await company_or_404(db, ent, company_id)
    data = _clean(body)
    _validate(data)
    row = AAEmployee(id=uuid.uuid4(), enterprise_id=ent, company_id=company.id, year=year, data=data)
    db.add(row)
    audit.record(db, label=audit.employee_label(data),
                 ent=ent, company_id=company.id, user=user, entity="employee", entity_id=row.id,
                 action="create", after={"year": year, **data})
    await db.commit()
    return await _out(db, company, year)


def _validate(data: dict) -> None:
    try:
        employee_from_json("check", data)
    except (ValueError, ArithmeticError) as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"Invalid value: {exc}") from None
    if not data.get("last_name"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Last name is required")


async def _employee_or_404(db: AsyncSession, ent: uuid.UUID, employee_id: uuid.UUID) -> AAEmployee:
    row = (await db.execute(select(AAEmployee).where(AAEmployee.id == employee_id,
                                                     AAEmployee.enterprise_id == ent))).scalar_one_or_none()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Employee not found")
    return row


@router.put("/employees/{employee_id}", response_model=EmployeesOut)
async def update_employee(employee_id: uuid.UUID, body: dict, user: User = Depends(get_current_user),
                          db: AsyncSession = Depends(get_db)):
    _require_manager(user)
    ent = scope_for(user)
    row = await _employee_or_404(db, ent, employee_id)
    data = _clean(body)
    _validate(data)
    audit.record(db, label=audit.employee_label(data),
                 ent=ent, company_id=row.company_id, user=user, entity="employee", entity_id=row.id,
                 action="update", before=row.data or {}, after=data)
    row.data = data
    await db.commit()
    company = await company_or_404(db, ent, row.company_id)
    return await _out(db, company, row.year)


@router.delete("/employees/{employee_id}", response_model=EmployeesOut)
async def delete_employee(employee_id: uuid.UUID, user: User = Depends(get_current_user),
                          db: AsyncSession = Depends(get_db)):
    _require_manager(user)
    ent = scope_for(user)
    row = await _employee_or_404(db, ent, employee_id)
    company_id, year = row.company_id, row.year
    audit.record(db, label=audit.employee_label(row.data),
                 ent=ent, company_id=company_id, user=user, entity="employee", entity_id=row.id,
                 action="delete", before={"year": year, **(row.data or {})}, after=None)
    await db.delete(row)
    await db.commit()
    company = await company_or_404(db, ent, company_id)
    return await _out(db, company, year)


# --- Downloads -----------------------------------------------------------------------------------------------------

_S1 = (("seq", "1", "Seq. No"), ("last", "2a", "Last Name"), ("first", "2b", "First Name"),
       ("middle", "2c", "Middle Name"), ("tin", "", "TIN"), ("3", "3", "Nationality/Resident (for foreigners only)"),
       ("4", "4", "Current Employment Status"), ("5a", "5a", "Period of Employment From (MM/DD)"),
       ("5b", "5b", "To (MM/DD)"), ("6", "6", "Reason of Separation, if applicable"),
       ("7a", "7a = (7f + 7j)", "Gross Compensation Income (present employer)"),
       ("7b", "7b", "13th Month Pay & Other Benefits"), ("7c", "7c", "De Minimis Benefits"),
       ("7d", "7d", "SSS, GSIS, PHIC, HDMF Contributions and Union Dues (employee share only)"),
       ("7e", "7e", "Salaries (P250,000 & below) & Other Forms of Compensation"),
       ("7f", "7f = (7b+7c+7d+7e)", "Total Non-Taxable/Exempt Compensation Income (present employer)"),
       ("7g", "7g", "Basic Salary (net of SSS, GSIS, PHIC, HDMF Contributions & Union Dues)"),
       ("7h", "7h", "13th Month Pay & Other Benefits (In Excess of Threshold)"),
       ("7i", "7i", "Salaries and Other Forms of Compensation"),
       ("7j", "7j = (7g+7h+7i)", "Total Taxable Compensation Income (present employer)"))
_S2 = (("seq", "1", "Seq. No."), ("last", "2a", "Last Name"), ("first", "2b", "First Name"),
       ("middle", "2c", "Middle Name"), ("tin", "", "TIN"), ("3", "3", "Current Employment Status"),
       ("4", "4", "Region No. Where Assigned"), ("5a", "5a", "From (MM/DD)"), ("5b", "5b", "To (MM/DD)"),
       ("6", "6", "Reason of Separation, if applicable"), ("7a", "7a=(7o+7r)", "Gross Compensation Income (present employer)"),
       ("7b", "7b", "Basic/SMW per Day"), ("7c", "7c", "Basic/SMW per Month"), ("7d", "7d", "Basic/SMW per Year"),
       ("7e", "7e", "Factor Used (No. of Days/Year)"),
       ("7f", "7f", "Basic/SMW (actual) (net of SSS, GSIS, PHIC, HDMF Contributions & Union Dues)"),
       ("7g", "7g", "Holiday Pay"), ("7h", "7h", "Overtime Pay"), ("7i", "7i", "Night Shift Differential"),
       ("7j", "7j", "Hazard Pay"), ("7k", "7k", "13th Month Pay & Other Benefits"), ("7l", "7l", "De Minimis Benefits"),
       ("7m", "7m", "SSS, GSIS, PHIC, HDMF Contributions & Union Dues (employee share only)"),
       ("7n", "7n", "Salaries and Other Forms of Compensation"),
       ("7o", "7o = (sum of 7f to 7n)", "Total Non-Taxable/Exempt Compensation (present employer)"),
       ("7p", "7p", "13th Month Pay & Other Benefits (In Excess of Threshold)"),
       ("7q", "7q", "Salaries and Other Forms of Compensation"),
       ("7r", "7r = (7p+7q)", "Total Taxable Compensation Income (present employer)"))
_PREV_1 = (("8", "8", "TIN (Previous Employer)"), ("9", "9", "Employment Status"), ("10a", "10a", "From (MM/DD)"),
           ("10b", "10b", "To (MM/DD)"), ("11", "11", "Reason of Separation, if applicable"),
           ("12a", "12a = (12f + 12j)", "Gross Compensation Income (previous employer)"),
           ("12b", "12b", "13th Month Pay & Other Benefits"), ("12c", "12c", "De Minimis Benefits"),
           ("12d", "12d", "SSS, GSIS, PHIC, HDMF Contributions & Union Dues (employee share only)"),
           ("12e", "12e", "Salaries (P250,000 and below) and Other Forms of Compensation"),
           ("12f", "12f=(12b+12c+12d+12e)", "Total Non-Taxable/Exempt Compensation Income (previous employer)"),
           ("12g", "12g", "Basic Salary (net of SSS, GSIS, PHIC, HDMF Contributions & Union Dues)"),
           ("12h", "12h", "13th Month Pay & Other Benefits"), ("12i", "12i", "Salaries and Other Forms of Compensation"),
           ("12j", "12j = (12g+12h+12i)", "Total Taxable Compensation Income (previous employer)"))
_PREV_2 = (("8", "8", "TIN (previous employer)"), ("9", "9", "Employment Status"), ("10a", "10a", "From (MM/DD)"),
           ("10b", "10b", "To (MM/DD)"), ("11", "11", "Reason of Separation, if applicable"),
           ("12a", "12a=(12k + 12n)", "Gross Compensation Income (previous employer)"),
           ("12b", "12b", "Basic/SMW (actual) (net of SSS, GSIS, PHIC, HDMF Contributions & Union Dues)"),
           ("12c", "12c", "Holiday Pay"), ("12d", "12d", "Overtime Pay"), ("12e", "12e", "Night Shift Differential"),
           ("12f", "12f", "Hazard Pay"), ("12g", "12g", "13th Month Pay & Other Benefits"),
           ("12h", "12h", "De Minimis Benefits"),
           ("12i", "12i", "SSS, GSIS, PHIC, HDMF Contributions & Union Dues (employee share only)"),
           ("12j", "12j", "Salaries and Other Forms of Compensation"),
           ("12k", "12k = (sum of 12b to 12j)", "Total Non-Taxable/Exempt Compensation (previous employer)"),
           ("12l", "12l", "13th Month Pay & Other Benefits"), ("12m", "12m", "Salaries and Other Forms of Compensation"),
           ("12n", "12n = (12l + 12m)", "Total Taxable Compensation (previous employer)"))
_TAIL = (("13", "13 = (7j + 12j)", "Total Taxable Compensation Income (Present & Previous Employer)"),
         ("14", "14", "Tax Due (January to December)"), ("15a", "15a", "Tax Withheld (Jan–Nov) Previous Employer"),
         ("15b", "15b", "Tax Withheld (Jan–Nov) Present Employer"), ("16", "16", "5% Tax Credit (PERA Act of 2008)"),
         ("17a", "17a = 14 - (15a + 15b) - 16", "Amount Withheld and Paid for in December or Last Salary"),
         ("17b", "17b = (15a + 15b + 16) - 14", "Over Withheld Tax Refunded to Employee"),
         ("18", "18 = (15a + 15b + 16 + 17a) OR (15a + 15b + 16 - 17b)",
          "Amount of Tax Withheld as Adjusted (to be reflected in BIR Form No. 2316 issued by the present employer)"),
         ("19", "19", "Substituted Filing? Yes/No"))
_HEAD = Font(bold=True, size=9)
_FILL = PatternFill("solid", start_color="FFE7E6E6")


def _sheet(wb: Workbook, title: str, heading: str, columns, rows: list[dict], company: AACompany, year: int) -> None:
    ws = wb.create_sheet(title)
    tin = company.tin
    ws.append([_safe(f"{company.registered_name} · TIN {tin[:3]}-{tin[3:6]}-{tin[6:9]}-{company.branch_code}")])
    ws.append([f"{heading} · for the year {year}"])
    ws["A1"].font = Font(bold=True, size=11)
    ws["A2"].font = Font(bold=True, size=10)
    ws.append([c[2] for c in columns])
    ws.append([c[1] for c in columns])
    for r in (3, 4):
        for col in range(1, len(columns) + 1):
            cell = ws.cell(r, col)
            cell.font, cell.fill = _HEAD, _FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[3].height = 60
    for row in rows:
        values = []
        for key, _, _ in columns:
            v = row.get(key, "")
            values.append(float(v) if isinstance(v, Decimal) else _safe(v))
        ws.append(values)
        if row.get("key") == "~total":
            for col in range(1, len(columns) + 1):
                ws.cell(ws.max_row, col).font = Font(bold=True)
    for i, (key, _, _) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 7 if key == "seq" else (18 if key in ("last", "first", "middle", "tin") else 15)
        if key not in ("seq", "last", "first", "middle", "tin", "3", "4", "5a", "5b", "6", "7e" if title == "Schedule 2" else "", "8", "9", "10a", "10b", "11", "19"):
            for r in range(5, ws.max_row + 1):
                ws.cell(r, i).number_format = "#,##0.00"
    ws.freeze_panes = "F5"


def alphalist_workbook(company: AACompany, year: int, tabs: dict, warnings: list[str]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Schedule 1", "Schedule 1 – Alphalist of Employees (Declared and Certified using BIR Form No. 2316)",
           _S1 + _PREV_1 + _TAIL, tabs["Schedule1"], company, year)
    _sheet(wb, "Schedule 2", "Schedule 2 – Alphalist of Minimum Wage Earners (Declared and Certified using BIR Form "
           "No. 2316)", _S2 + _PREV_2 + _TAIL, tabs["Schedule2"], company, year)
    _sheet(wb, "1604-C remittances", "BIR Form 1604-C Part II – Summary of Remittances (from the 1601-C returns)",
           (("month", "", "Month"), ("remitted_on", "", "Date of Remittance"), ("withheld", "", "Taxes Withheld"),
            ("adjustment", "", "Adjustment"), ("penalties", "", "Penalties"), ("total", "", "Total Amount Remitted"),
            ("source", "", "Source")), tabs["Remittances"], company, year)
    _sheet(wb, "Checks", "Alphalist vs 1601-C returns",
           (("check", "", "Check"), ("alphalist", "", "Alphalist"), ("returns", "", "1601-C returns"),
            ("difference", "", "Difference")), tabs["Checks"], company, year)
    ws = wb.create_sheet("Warnings")
    ws.append(["Warnings"])
    for w in warnings or ["None."]:
        ws.append([_safe(w)])
    ws.column_dimensions["A"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _filename(company: AACompany, what: str, year: int, ext: str) -> str:
    name = "".join(ch for ch in company.registered_name if ch.isalnum() or ch in " .-")[:60].strip()
    return f"{name} {what} {year}.{ext}"


@router.get("/companies/{company_id}/employees/{year}/alphalist.xlsx")
async def download_alphalist(company_id: uuid.UUID, year: int, user: User = Depends(get_current_user),
                             db: AsyncSession = Depends(get_db)):
    company = await company_or_404(db, scope_for(user), company_id)
    _, tabs, _, warnings = await _build(db, company, year)
    data = await asyncio.get_running_loop().run_in_executor(None, alphalist_workbook, company, year, tabs, warnings)
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_filename(company, "1604-C alphalist", year, "xlsx")}"'})


async def _computed_2316(db: AsyncSession, user: User, company_id: uuid.UUID, year: int):
    company = await company_or_404(db, scope_for(user), company_id)
    _, _, computed, _ = await _build(db, company, year)
    if not computed:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "No employees for that year")
    employer = {"tin": company.tin + company.branch_code, "name": company.registered_name,
                "address": company.address, "zip": company.zip_code}
    return company, computed, employer


def _2316_label(c) -> str:
    return f"{c.employee.last_name}, {c.employee.first_name}".strip(", ") or "(no name)"


def _edited_2316(c, values: dict[str, str]):
    try:
        return apply_2316(c, values)
    except ValueError as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(e))


async def _2316_zip(db: AsyncSession, user: User, company_id: uuid.UUID, year: int,
                    edits: dict[str, dict[str, str]]) -> Response:
    company, computed, employer = await _computed_2316(db, user, company_id, year)
    ids = {c.employee.id for c in computed}
    if any(key not in ids for key in edits):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Some edited employees are no longer in this year")
    computed = [_edited_2316(c, edits[c.employee.id]) if c.employee.id in edits else c for c in computed]
    data = await asyncio.get_running_loop().run_in_executor(None, build_2316_zip, computed, employer, year)
    return Response(data, media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{_filename(company, "2316s", year, "zip")}"'})


@router.get("/companies/{company_id}/employees/{year}/2316.zip")
async def download_2316s(company_id: uuid.UUID, year: int, user: User = Depends(get_current_user),
                         db: AsyncSession = Depends(get_db)):
    return await _2316_zip(db, user, company_id, year, {})


@router.post("/companies/{company_id}/employees/{year}/2316.zip")
async def download_edited_2316s(company_id: uuid.UUID, year: int, body: CertificatesIn,
                                user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    """The zip with the preview's edits (they only change these files, not the employee records)."""
    return await _2316_zip(db, user, company_id, year, body.edits)


@router.get("/companies/{company_id}/employees/{year}/2316/certificates", response_model=list[CertificateOut])
async def list_2316s(company_id: uuid.UUID, year: int, user: User = Depends(get_current_user),
                     db: AsyncSession = Depends(get_db)):
    _, computed, _ = await _computed_2316(db, user, company_id, year)
    return [CertificateOut(id=c.employee.id, label=_2316_label(c), fields=[PdfFieldOut(**f) for f in fields_2316(c)])
            for c in computed]


@router.post("/companies/{company_id}/employees/{year}/2316/preview")
async def preview_2316(company_id: uuid.UUID, year: int, body: CertificatePreviewIn,
                       user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    import fitz

    _, computed, employer = await _computed_2316(db, user, company_id, year)
    c = next((c for c in computed if c.employee.id == body.id), None)
    if c is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Employee not found")
    c = _edited_2316(c, body.values)

    def render() -> bytes:
        doc = fitz.open()
        fill_2316(doc, c, employer)
        return doc.tobytes(garbage=3, deflate=True)

    data = await asyncio.get_running_loop().run_in_executor(None, render)
    return Response(content=data, media_type="application/pdf",
                    headers={"Content-Disposition": 'inline; filename="2316 preview.pdf"', "Cache-Control": "no-store"})


@router.get("/alphalist-status", response_model=list[StatusRow])
async def alphalist_status(year: int = Query(..., ge=2000, le=2100), user: User = Depends(get_current_user),
                           db: AsyncSession = Depends(get_db)):
    """Every client with employees: is its 1604-C alphalist ready for the year?"""
    ent = scope_for(user)
    companies = (await db.execute(select(AACompany).where(
        AACompany.enterprise_id == ent, AACompany.has_employees.is_(True)).order_by(AACompany.registered_name)
    )).scalars().all()
    counts = dict((await db.execute(select(AAEmployee.company_id, func.count()).where(
        AAEmployee.enterprise_id == ent, AAEmployee.year == year).group_by(AAEmployee.company_id))).all())
    out = []
    for c in companies:
        n = counts.get(c.id, 0)
        issues = 0
        if n:
            _, _, _, warnings = await _build(db, c, year)
            issues = len(warnings)
        out.append(StatusRow(company_id=c.id, company=c.registered_name, employees=n, issues=issues,
                             status="no_payroll" if not n else ("needs_review" if issues else "ready")))
    return out
