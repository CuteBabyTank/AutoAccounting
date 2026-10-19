"""Spreadsheet importers: opening trial balance, opening open items (A/R / A/P) and journal entries.

Each importer is a stateless preview → confirm pair: the browser uploads the file with the person's choices (columns,
date format, account mapping fixes, where an out-of-balance difference goes); the preview shows what would post; the
confirm re-reads the same file with the same choices and posts only when it is ready. Writes go through the existing
paths: the opening TB is the company's opening journal entry, journals are manual entries checked like the JE form,
opening items are rows the books, aging and bank matcher pick up.
"""

import asyncio
import csv
import io
import json
import re
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

from fastapi import APIRouter, Depends, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_current_user, get_db
from app.models.accounting import (
    AABankLine, AABankReconciliation, AACompany, AAImportBatch, AAJournalEntry, AAOpeningItem, AAPeriod,
)
from app.models.user import User
from app.routers.accounting_books import ENTRY_FIELDS, entry_label, entry_problem, lines_json, new_entry_row
from app.routers.accounting_scope import MANAGER_ROLES, company_or_404, scope_for
from app.schemas.accounting_books import JournalEntryIn
from app.schemas.accounting_imports import (
    AccountChoiceOut, DateFormatOut, ImportBatchesOut, ImportBatchOut, ImportKind, ImportOptions, ImportPreviewOut,
    ImportResultOut, ImportUndoOut, JePreviewOut,
    OpeningItemOut, OpeningItemPatch, OpeningItemsOut, OpenItemPreviewOut, RowIssueOut, TbLineOut,
)
from app.services.accounting import audit
from app.services.accounting.books import bank
from app.services.accounting.books.chart import Chart
from app.services.accounting.books.loader import active_closes, ensure_chart, opening_items, touch_books
from app.services.accounting.compute.bank_books import POSTING_STATUSES
from app.services.accounting.compute.money import ZERO, c2
from app.services.accounting.import_undo import (
    KIND_LABELS, UndoFacts, bank_touches, batch_status, closed_hits, entries_plan, entry_state, filed_hits,
    fingerprint, opening_bank_delta, opening_plan, reconciled_hits, undo_problem,
)
from app.services.accounting.import_rows import (
    JE_FIELDS, JE_SYNONYMS, OPEN_FIELDS, OPEN_SYNONYMS, TB_FIELDS, TB_SYNONYMS, TEMPLATES, AccountChoice,
    build_journals, build_open_items, build_tb, control_tie, document_keys, item_key,
)
from app.services.accounting.importing import DATE_LABELS, AccountRef, suggest_columns
from app.services.accounting.opening_items import control_balances
from app.services.accounting.payroll_import import parse_sheet
from app.services.spreadsheet_safe import csv_safe

router = APIRouter()

MAX_IMPORT_BYTES = 5 * 1024 * 1024
PREVIEW_ROWS = 300
_FIELDS = {"opening-tb": (TB_FIELDS, TB_SYNONYMS), "open-items": (OPEN_FIELDS, OPEN_SYNONYMS),
           "journal-entries": (JE_FIELDS, JE_SYNONYMS)}
_ITEM_FIELDS = ("kind", "contact_name", "contact_tin", "doc_number", "doc_date", "due_date", "terms_days", "amount",
                "paid_on", "paid_via", "paid_account")


def _require_manager(user: User) -> None:
    if user.role not in MANAGER_ROLES:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Only owners and managers can import into the books")


def parse_options(raw: str | None) -> ImportOptions:
    try:
        data = json.loads(raw) if raw else {}
        if not isinstance(data, dict):
            raise ValueError
        return ImportOptions.model_validate(data)
    except (ValueError, ValidationError):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid import options") from None


async def _read(file: UploadFile) -> tuple[list[str], list[list[str]]]:
    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "The file is larger than 5 MB")
    try:
        return await asyncio.get_running_loop().run_in_executor(None, parse_sheet, data, file.filename or "")
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from None
    except Exception:  # a corrupt workbook
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Couldn't read that file as Excel or CSV") from None


def chart_refs(chart: Chart) -> list[AccountRef]:
    return [AccountRef(a.code, a.name, a.active, a.locked) for a in chart.accounts]


def chosen_columns(kind: str, headers: list[str], opts: ImportOptions) -> dict[str, str]:
    fields, synonyms = _FIELDS[kind]
    if opts.columns is None:
        return suggest_columns(headers, synonyms)
    return {f: h for f, h in opts.columns.items() if f in fields and h in headers}


def _account_out(a: AccountChoice) -> AccountChoiceOut:
    return AccountChoiceOut(key=a.key, code_in=a.code_in, name_in=a.name_in, match_code=a.match.code,
                            how=a.match.how, score=a.match.score, alternatives=a.match.alternatives, chosen=a.chosen,
                            problem=a.problem, rows=a.rows)


def _formats(candidates: list[str]) -> list[DateFormatOut]:
    return [DateFormatOut(value=c, label=DATE_LABELS.get(c, c)) for c in candidates]


async def _opening_entry(db: AsyncSession, company: AACompany) -> AAJournalEntry | None:
    return (await db.execute(select(AAJournalEntry).where(
        AAJournalEntry.enterprise_id == company.enterprise_id, AAJournalEntry.company_id == company.id,
        AAJournalEntry.kind == "opening", AAJournalEntry.deleted_at.is_(None)).limit(1))).scalar_one_or_none()


async def _control(db: AsyncSession, company: AACompany, chart: Chart) -> dict[str, Decimal]:
    entry = await _opening_entry(db, company)
    return control_balances(entry.lines if entry else [], chart)


async def _uploaded_keys(db: AsyncSession, company: AACompany) -> set[tuple[str, str, str]]:
    """The sales / bills already uploaded (they age on their own, so an opening item for one would count twice)."""
    from app.services.accounting.books.subledger import company_documents, counterparty, doc_kind

    out = []
    for d in await company_documents(db, company):
        kind = doc_kind(d)
        if kind is not None:
            tin, name, _ = counterparty(d)
            out.append((kind, name, tin, d.doc_number or ""))
    return document_keys(out)


async def _posted_keys(db: AsyncSession, company: AACompany) -> set[tuple[date, str, Decimal]]:
    """(date, reference, total debit) of the company's journal entries, to spot a file imported twice."""
    rows = (await db.execute(select(AAJournalEntry.entry_date, AAJournalEntry.reference, AAJournalEntry.lines).where(
        AAJournalEntry.enterprise_id == company.enterprise_id, AAJournalEntry.company_id == company.id,
        AAJournalEntry.kind == "manual", AAJournalEntry.deleted_at.is_(None)))).all()
    return {(on, (ref or "").strip(), sum((c2(l.get("debit") or 0) for l in lines or []), ZERO))
            for on, ref, lines in rows}


def _totals(items) -> dict[str, Decimal]:
    out = {"ar": ZERO, "ap": ZERO}
    for i in items:
        out[i.kind] = c2(out.get(i.kind, ZERO) + c2(i.amount))
    return out


class _Built:
    def __init__(self, preview: ImportPreviewOut, result, chart: Chart):
        self.preview, self.result, self.chart = preview, result, chart


async def _build(db: AsyncSession, company: AACompany, kind: str, file: UploadFile, opts: ImportOptions) -> _Built:
    if company.books_start is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Set the books start date on the company first")
    headers, rows = await _read(file)
    chart = await ensure_chart(db, company)
    columns = chosen_columns(kind, headers, opts)
    common = dict(kind=kind, filename=(file.filename or "")[:120], headers=headers, columns=columns,
                  fields=_FIELDS[kind][0], row_count=len(rows), books_start=company.books_start)
    refs = chart_refs(chart)
    if kind == "opening-tb":
        r = build_tb(headers, rows, columns, refs, opts.accounts, date_format=opts.date_format,
                     books_start=company.books_start, difference_account=opts.difference_account or None)
        preview = ImportPreviewOut(
            **common, date_format=r.date_format, date_formats=_formats(r.date_candidates),
            accounts=[_account_out(a) for a in r.accounts], errors=[RowIssueOut(**vars(e)) for e in r.errors],
            warnings=r.warnings, ready=r.ready,
            tb_lines=[TbLineOut(account_code=c, account_name=chart.name(c), debit=d, credit=k) for c, d, k in r.lines],
            debit=r.debit, credit=r.credit, difference=r.difference, raw_difference=r.raw_difference, as_of=r.as_of,
            has_opening=await _opening_entry(db, company) is not None)
        closes = await active_closes(db, company.id)
        first = next((c for c in closes if c.month_start <= company.books_start <= c.month_end), None)
        if first is not None:
            preview.warnings.append(
                f"{first.month_start:%B %Y} is closed: a change to the opening balances posts as a correcting entry on "
                f"the first open day and the closed pages keep the old figures. Reopen the month first to change the "
                "opening balances themselves.")
        if r.lines:
            ar_ap = control_balances([{"account_code": c, "debit": d, "credit": k} for c, d, k in r.lines], chart)
            existing = _totals(await opening_items(db, company))
            if any(existing.values()):
                preview.warnings += control_tie(existing, ar_ap)
        return _Built(preview, r, chart)
    if kind == "open-items":
        current = await opening_items(db, company)
        r = build_open_items(headers, rows, columns, company.books_start, default_kind=opts.default_kind,
                             date_format=opts.date_format,
                             existing={item_key(i.kind, i.contact_name, i.doc_number) for i in current},
                             documents=await _uploaded_keys(db, company))
        control = await _control(db, company, chart)
        after = _totals(current)
        for k, v in r.totals.items():
            after[k] = c2(after.get(k, ZERO) + v)
        warnings = list(r.warnings) + (control_tie(after, control) if r.items else [])
        preview = ImportPreviewOut(
            **common, date_format=r.date_format, date_formats=_formats(r.date_candidates),
            errors=[RowIssueOut(**vars(e)) for e in r.errors], warnings=warnings, ready=r.ready,
            items=[OpenItemPreviewOut(**{k: v for k, v in vars(i).items()}) for i in r.items[:PREVIEW_ROWS]],
            item_count=len(r.items), totals=r.totals, control=control, skipped=r.skipped + r.documented)
        return _Built(preview, r, chart)
    closes = await active_closes(db, company.id)
    last_closed = closes[-1].month_end if closes else None
    r = build_journals(headers, rows, columns, refs, opts.accounts, books_start=company.books_start,
                       date_format=opts.date_format, last_closed=last_closed, posted=await _posted_keys(db, company))
    preview = ImportPreviewOut(
        **common, date_format=r.date_format, date_formats=_formats(r.date_candidates),
        accounts=[_account_out(a) for a in r.accounts], errors=[RowIssueOut(**vars(e)) for e in r.errors],
        warnings=r.warnings, ready=r.ready, last_closed=last_closed,
        entries=[JePreviewOut(row=e.row, entry_date=e.entry_date, posts_on=e.posts_on, reference=e.reference,
                              particulars=e.particulars, is_adjusting=e.is_adjusting, debit=e.debit, credit=e.credit,
                              lines=len(e.lines), errors=e.errors) for e in r.entries[:PREVIEW_ROWS]])
    return _Built(preview, r, chart)


@router.post("/companies/{company_id}/imports/{kind}/preview", response_model=ImportPreviewOut)
async def preview_import(company_id: uuid.UUID, kind: ImportKind, file: UploadFile, options: str = Form(""),
                         user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    _require_manager(user)
    company = await company_or_404(db, scope_for(user), company_id)
    built = await _build(db, company, kind, file, parse_options(options))
    return built.preview


def _entry_body(**kw) -> JournalEntryIn:
    try:
        return JournalEntryIn(**kw)
    except ValidationError as exc:
        message = "; ".join(str(e.get("msg", "")).removeprefix("Value error, ") for e in exc.errors()[:3])
        raise HTTPException(status.HTTP_400_BAD_REQUEST, message or "Invalid journal entry") from None


@router.post("/companies/{company_id}/imports/{kind}", response_model=ImportResultOut)
async def confirm_import(company_id: uuid.UUID, kind: ImportKind, file: UploadFile, options: str = Form(""),
                         user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    _require_manager(user)
    ent = scope_for(user)
    company = await company_or_404(db, ent, company_id)
    await _lock_company(db, company)   # an import and an undo of the same company don't interleave
    built = await _build(db, company, kind, file, parse_options(options))
    if not built.preview.ready:
        first = built.preview.errors[0].message if built.preview.errors else "Fix the problems shown first"
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Not ready to import: {first}")
    name = built.preview.filename
    out = ImportResultOut(kind=kind, warnings=built.preview.warnings)
    batch = AAImportBatch(id=uuid.uuid4(), enterprise_id=ent, company_id=company.id, kind=kind, filename=name,
                          created_by=user.id)
    db.add(batch)
    await db.flush()   # the rows below reference it
    written: dict[str, str] = {}
    if kind == "opening-tb":
        body = _entry_body(kind="opening", book="GJ", entry_date=company.books_start, reference="",
                           particulars=f"Opening balances (imported from {name})"[:255] if name else "Opening balances",
                           lines=[{"account_code": c, "debit": d, "credit": k} for c, d, k in built.result.lines])
        row = await _opening_entry(db, company)
        kept = {str(x.get("account_code", "")) for x in (row.lines or [])} if row else set()
        problem = entry_problem(built.chart, body, company.books_start, kept)
        if problem:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, problem)
        batch.restore = {"opening": entry_state(row) if row is not None else None}
        if row is None:
            row = new_entry_row(ent, company, body, user)
            row.import_batch_id = batch.id
            db.add(row)
            audit.record(db, label=entry_label(row), ent=ent, company_id=company.id, user=user,
                         entity="journal_entry", entity_id=row.id, action="import",
                         after=audit.snapshot(row, ENTRY_FIELDS), reason=f"Imported from {name}"[:500])
            out.created = 1
        else:
            before = audit.snapshot(row, ENTRY_FIELDS)
            row.lines, row.particulars = lines_json(body), body.particulars
            row.import_batch_id = batch.id
            audit.record(db, label=entry_label(row), ent=ent, company_id=company.id, user=user,
                         entity="journal_entry", entity_id=row.id, action="import", before=before,
                         after=audit.snapshot(row, ENTRY_FIELDS), reason=f"Imported from {name}"[:500], always=True)
            out.updated = 1
        written[str(row.id)] = fingerprint(entry_state(row))
        batch.row_count = len(row.lines or [])
        touch_books(company)
        await db.flush()
        await bank.refresh_opening_flag(db, company)
        if (company.review_flags or {}).get("opening_balance_missing"):
            out.warnings.append("The opening cash in bank still doesn't match the first bank statement's opening "
                                "balance (see Bank).")
    elif kind == "open-items":
        for i in built.result.items:
            db.add(AAOpeningItem(enterprise_id=ent, company_id=company.id, kind=i.kind, contact_name=i.contact_name,
                                 contact_tin=i.contact_tin, doc_number=i.doc_number, doc_date=i.doc_date,
                                 due_date=i.due_date, terms_days=i.terms_days, amount=i.amount, created_by=user.id,
                                 import_batch_id=batch.id))
        totals = built.result.totals
        audit.record(db, label=f"Opening open items · {name}"[:200], ent=ent, company_id=company.id, user=user,
                     entity="opening_item", entity_id=str(company.id), action="import", before={},
                     after={"items": len(built.result.items), "receivables": totals.get("ar", ZERO),
                            "payables": totals.get("ap", ZERO), "file": name}, always=True)
        out.created, out.skipped = len(built.result.items), built.result.skipped + built.result.documented
        batch.row_count = out.created
        await db.flush()
        await bank.rematch(db, company)
    else:
        for draft in built.result.entries:
            body = _entry_body(kind="manual", book="GJ", entry_date=draft.entry_date, reference=draft.reference,
                               particulars=draft.particulars, is_adjusting=draft.is_adjusting, lines=draft.lines)
            problem = entry_problem(built.chart, body, company.books_start)
            if problem:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Row {draft.row}: {problem}")
            row = new_entry_row(ent, company, body, user)
            row.import_batch_id = batch.id
            db.add(row)
            audit.record(db, label=entry_label(row), ent=ent, company_id=company.id, user=user,
                         entity="journal_entry", entity_id=row.id, action="import",
                         after=audit.snapshot(row, ENTRY_FIELDS), reason=f"Imported from {name}"[:500])
            written[str(row.id)] = fingerprint(entry_state(row))
            out.created += 1
        batch.row_count = out.created
        touch_books(company)
    if batch.row_count or kind == "opening-tb":
        batch.written = written
        out.batch_id = batch.id
        audit.record(db, label=_batch_label(batch), ent=ent, company_id=company.id, user=user, entity="import_batch",
                     entity_id=batch.id, action="import", before=None,
                     after={"kind": kind, "file": name, "rows": batch.row_count}, reason=f"Imported from {name}"[:500])
    else:   # nothing written (every open item was already there): nothing to undo
        await db.delete(batch)
    await db.commit()
    return out


async def _lock_company(db: AsyncSession, company: AACompany) -> None:
    await db.execute(select(AACompany.id).where(AACompany.id == company.id,
                                                AACompany.enterprise_id == company.enterprise_id).with_for_update())


def _batch_label(batch: AAImportBatch) -> str:
    what = KIND_LABELS.get(batch.kind, batch.kind).capitalize()
    return " · ".join(p for p in (f"{what} import", batch.filename) if p)[:200]


# --- import batches (recent imports, undo) --------------------------------------------------------------------------

BATCH_LIST_LIMIT = 50


async def _emails(db: AsyncSession, ids: set) -> dict:
    ids = {i for i in ids if i is not None}
    if not ids:
        return {}
    return dict((await db.execute(select(User.id, User.email).where(User.id.in_(ids)))).all())


def _batch_out(batch: AAImportBatch, emails: dict, opening_tag: str | None) -> ImportBatchOut:
    return ImportBatchOut(id=batch.id, kind=batch.kind, filename=batch.filename, row_count=batch.row_count,
                          created_by=emails.get(batch.created_by, ""), created_at=batch.created_at,
                          undone_at=batch.undone_at, undone_by=emails.get(batch.undone_by, ""),
                          status=batch_status(batch.kind, str(batch.id), batch.undone_at is not None, opening_tag))


async def _opening_tag(db: AsyncSession, company: AACompany) -> str | None:
    entry = await _opening_entry(db, company)
    return str(entry.import_batch_id) if entry is not None and entry.import_batch_id else None


@router.get("/companies/{company_id}/imports/batches", response_model=ImportBatchesOut)
async def list_import_batches(company_id: uuid.UUID, user: User = Depends(get_current_user),
                              db: AsyncSession = Depends(get_db)):
    ent = scope_for(user)
    company = await company_or_404(db, ent, company_id)
    rows = (await db.execute(select(AAImportBatch).where(
        AAImportBatch.enterprise_id == ent, AAImportBatch.company_id == company.id)
        .order_by(AAImportBatch.created_at.desc()).limit(BATCH_LIST_LIMIT))).scalars().all()
    emails = await _emails(db, {r.created_by for r in rows} | {r.undone_by for r in rows})
    tag = await _opening_tag(db, company)
    return ImportBatchesOut(batches=[_batch_out(r, emails, tag) for r in rows], can_undo=user.role in MANAGER_ROLES)


async def _batch_or_404(db: AsyncSession, ent: uuid.UUID, batch_id: uuid.UUID, lock: bool = False) -> AAImportBatch:
    q = select(AAImportBatch).where(AAImportBatch.id == batch_id, AAImportBatch.enterprise_id == ent)
    if lock:
        q = q.with_for_update().execution_options(populate_existing=True)
    row = (await db.execute(q)).scalar_one_or_none()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Import not found")
    return row


async def _common_facts(db: AsyncSession, company: AACompany):
    """(closed months, filed quarters, live reconciliations, bank account codes) of the company."""
    closes = [(c.month_start, c.month_end) for c in await active_closes(db, company.id)]
    periods = (await db.execute(select(AAPeriod.year, AAPeriod.quarter, AAPeriod.filed_at, AAPeriod.filed_forms).where(
        AAPeriod.enterprise_id == company.enterprise_id, AAPeriod.company_id == company.id))).all()
    filed = {(y, q) for y, q, at, forms in periods if at is not None or forms}
    recs = [tuple(r) for r in (await db.execute(select(
        AABankReconciliation.account_code, AABankReconciliation.month_start, AABankReconciliation.month_end).where(
        AABankReconciliation.enterprise_id == company.enterprise_id, AABankReconciliation.company_id == company.id,
        AABankReconciliation.voided_at.is_(None)))).all()]
    codes = {a.code for a in await bank.bank_accounts(db, company.id)}
    return closes, filed, recs, codes


async def _lines_matched_to(db: AsyncSession, company: AACompany, kind: str, refs: list[str]) -> list[AABankLine]:
    """Live statement lines whose posting match (auto or confirmed) points at these rows."""
    if not refs:
        return []
    return list((await db.execute(select(AABankLine).where(
        AABankLine.enterprise_id == company.enterprise_id, AABankLine.company_id == company.id,
        AABankLine.match_kind == kind, AABankLine.matched_ref.in_(refs),
        AABankLine.match_status.in_(POSTING_STATUSES), bank.live_line()))).scalars().all())


@router.post("/import-batches/{batch_id}/undo", response_model=ImportUndoOut)
async def undo_import(batch_id: uuid.UUID, user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    """Undo a whole import: remove the rows it wrote (the opening TB puts back the opening entry it replaced).

    Refused (409) when the import is already undone, posts into a closed month or filed quarter, is inside a bank
    reconciliation, or later work depends on its rows (a confirmed bank match, a payment, a write-off, an edit).
    """
    if user.role not in MANAGER_ROLES:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Only owners and managers can undo an import")
    ent = scope_for(user)
    batch = await _batch_or_404(db, ent, batch_id)
    company = await company_or_404(db, ent, batch.company_id)
    await _lock_company(db, company)
    batch = await _batch_or_404(db, ent, batch_id, lock=True)   # re-read under the lock: undone by someone else?
    closes, filed, recs, codes = await _common_facts(db, company)
    now = datetime.now(timezone.utc)
    reason = f"Undo of the import of {batch.filename}"[:500] if batch.filename else "Undo of an import"
    out_removed, restored = 0, False
    dates: list[date] = []
    touches: list[tuple[str, date]] = []
    extra: dict = {}

    if batch.kind == "opening-tb":
        current = await _opening_entry(db, company)
        plan = opening_plan(str(batch.id), batch.restore or {}, batch.written or {}, current)
        prior_lines = (plan.restore or {}).get("lines") or []
        if current is not None and company.books_start is not None:
            dates = [company.books_start]
            touches = [(c, company.books_start) for c in opening_bank_delta(current.lines or [], prior_lines, codes)]
        facts = UndoFacts(kind=batch.kind, undone=batch.undone_at is not None, fye=company.fiscal_year_end_month,
                          opening=plan.problem)
        extra = dict(current=current, plan=plan)
    elif batch.kind == "journal-entries":
        live = list((await db.execute(select(AAJournalEntry).where(
            AAJournalEntry.enterprise_id == ent, AAJournalEntry.company_id == company.id,
            AAJournalEntry.import_batch_id == batch.id, AAJournalEntry.deleted_at.is_(None))
            .order_by(AAJournalEntry.entry_date))).scalars().all())
        _, edited = entries_plan(batch.written or {}, live)
        by_id = {str(e.id): e for e in live}
        lines = await _lines_matched_to(db, company, "journal", list(by_id))
        dates = [e.entry_date for e in live]
        touches = bank_touches(live, codes) + [(l.account_code, l.line_date) for l in lines]
        facts = UndoFacts(kind=batch.kind, undone=batch.undone_at is not None, fye=company.fiscal_year_end_month,
                          confirmed=[entry_label(by_id[l.matched_ref]) for l in lines if l.match_status == "confirmed"],
                          edited=[entry_label(by_id[i]) for i in edited])
        extra = dict(live=live)
    else:
        from app.services.accounting.books.subledger import writeoff_rows

        items = list((await db.execute(select(AAOpeningItem).where(
            AAOpeningItem.enterprise_id == ent, AAOpeningItem.company_id == company.id,
            AAOpeningItem.import_batch_id == batch.id, AAOpeningItem.deleted_at.is_(None)))).scalars().all())
        by_id = {str(i.id): i for i in items}
        lines = await _lines_matched_to(db, company, "document", list(by_id))
        off = {str(x.get("id")) for _, _, xs in await writeoff_rows(db, company) for x in xs}
        dates = [l.line_date for l in lines]   # the items post nothing themselves; their matched payments do
        touches = [(l.account_code, l.line_date) for l in lines]
        facts = UndoFacts(kind=batch.kind, undone=batch.undone_at is not None, fye=company.fiscal_year_end_month,
                          confirmed=[_item_label(by_id[l.matched_ref]) for l in lines if l.match_status == "confirmed"],
                          paid=[_item_label(i) for i in items if i.paid_on is not None],
                          written_off=[_item_label(i) for k, i in by_id.items() if k in off])
        extra = dict(items=items)

    if not facts.undone:
        facts = UndoFacts(**{**vars(facts), "closed": closed_hits(dates, closes),
                             "filed": filed_hits(dates, filed, company.fiscal_year_end_month),
                             "reconciled": reconciled_hits(touches, recs)})
    problem = undo_problem(facts)
    if problem:
        raise HTTPException(status.HTTP_409_CONFLICT, problem)

    if batch.kind == "opening-tb":
        current, plan = extra["current"], extra["plan"]
        before = audit.snapshot(current, ENTRY_FIELDS)
        if plan.action == "delete":
            current.deleted_at = now
            after = None
        else:
            prior = plan.restore
            current.kind, current.book = prior["kind"], prior["book"]
            current.entry_date = date.fromisoformat(prior["entry_date"])
            current.reference, current.particulars = prior["reference"], prior["particulars"]
            current.lines, current.is_adjusting = prior["lines"], bool(prior["is_adjusting"])
            current.import_batch_id = uuid.UUID(prior["import_batch_id"]) if prior.get("import_batch_id") else None
            after = audit.snapshot(current, ENTRY_FIELDS)
            restored = True
        audit.record(db, label=entry_label(current), ent=ent, company_id=company.id, user=user,
                     entity="journal_entry", entity_id=current.id, action="undo_import", before=before, after=after,
                     reason=reason, always=True)
        out_removed = 0 if restored else 1
    elif batch.kind == "journal-entries":
        for e in extra["live"]:
            e.deleted_at = now
            audit.record(db, label=entry_label(e), ent=ent, company_id=company.id, user=user, entity="journal_entry",
                         entity_id=e.id, action="undo_import", before=audit.snapshot(e, ENTRY_FIELDS), after=None,
                         reason=reason)
        out_removed = len(extra["live"])
    else:
        items = extra["items"]
        for i in items:
            i.deleted_at = now
        totals = _totals(items)
        if items:
            audit.record(db, label=f"Opening open items · {batch.filename}"[:200], ent=ent, company_id=company.id,
                         user=user, entity="opening_item", entity_id=str(company.id), action="undo_import",
                         before={"items": len(items), "receivables": totals.get("ar", ZERO),
                                 "payables": totals.get("ap", ZERO), "file": batch.filename}, after=None,
                         reason=reason)
        out_removed = len(items)

    batch.undone_at, batch.undone_by = now, user.id
    audit.record(db, label=_batch_label(batch), ent=ent, company_id=company.id, user=user, entity="import_batch",
                 entity_id=batch.id, action="undo_import", before={"status": "active"},
                 after={"status": "undone", "removed": out_removed, "restored": restored}, reason=reason)
    touch_books(company)
    await db.flush()
    for account, on in touches:   # none hold today (the guard refuses), kept so a changed line never leaves one
        await bank.void_reconciliations(db, company.id, account, on)
    await bank.rematch(db, company)
    await bank.refresh_opening_flag(db, company)
    warnings = []
    if (company.review_flags or {}).get("opening_balance_missing"):
        warnings.append("The opening cash in bank doesn't match the first bank statement's opening balance (see Bank).")
    emails = await _emails(db, {batch.created_by, batch.undone_by})
    out = ImportUndoOut(batch=_batch_out(batch, emails, await _opening_tag(db, company)), removed=out_removed,
                        restored=restored, warnings=warnings)
    await db.commit()
    return out


# --- templates -----------------------------------------------------------------------------------------------------

def _cell(value: str):
    """Numbers stay numbers (so a template's negative amount isn't turned into text); text is formula-guarded."""
    return Decimal(value) if re.fullmatch(r"-?\d+\.\d+", value or "") else csv_safe(value)


def template_bytes(kind: str, fmt: str) -> bytes:
    headers, rows = TEMPLATES[kind]
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([csv_safe(h) for h in headers])
        for r in rows:
            writer.writerow([_cell(c) for c in r])
        return buf.getvalue().encode("utf-8-sig")
    wb = Workbook()
    ws = wb.active
    ws.title = "Import"
    ws.append([csv_safe(h) for h in headers])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in (_cell(c) for c in r)])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(h) + 4)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@router.get("/imports/templates/{kind}")
async def import_template(kind: ImportKind, format: str = Query("xlsx", pattern="^(csv|xlsx)$"),
                          user: User = Depends(get_current_user)):
    data = template_bytes(kind, format)
    media = "text/csv" if format == "csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(data, media_type=media,
                    headers={"Content-Disposition": f'attachment; filename="{kind}-template.{format}"'})


# --- opening open items ---------------------------------------------------------------------------------------------

@router.get("/companies/{company_id}/opening-items", response_model=OpeningItemsOut)
async def list_opening_items(company_id: uuid.UUID, user: User = Depends(get_current_user),
                             db: AsyncSession = Depends(get_db)):
    from app.services.accounting.books.subledger import load_subledger

    company = await company_or_404(db, scope_for(user), company_id)
    items = await opening_items(db, company)
    chart = await ensure_chart(db, company)
    open_by_id: dict[str, Decimal] = {}
    if items:
        sub = await load_subledger(db, company, max(date.today(), company.books_start))
        open_by_id = {i.doc.id: i.open for i in sub.items if i.doc.opening}
    control = await _control(db, company, chart)
    totals = _totals(items)
    await db.commit()
    out = []
    for i in items:
        row = OpeningItemOut.model_validate(i)
        row.open = open_by_id.get(str(i.id), ZERO)
        out.append(row)
    return OpeningItemsOut(items=out, totals=totals, control=control,
                           warnings=control_tie(totals, control) if items else [], books_start=company.books_start)


async def _item_or_404(db: AsyncSession, ent: uuid.UUID, item_id: uuid.UUID) -> AAOpeningItem:
    row = (await db.execute(select(AAOpeningItem).where(
        AAOpeningItem.id == item_id, AAOpeningItem.enterprise_id == ent, AAOpeningItem.deleted_at.is_(None))
    )).scalar_one_or_none()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Opening item not found")
    return row


def _item_label(row: AAOpeningItem) -> str:
    return " · ".join(p for p in ("A/R" if row.kind == "ar" else "A/P", row.doc_number, row.contact_name) if p)[:200]


@router.patch("/opening-items/{item_id}", response_model=OpeningItemOut)
async def patch_opening_item(item_id: uuid.UUID, body: OpeningItemPatch, user: User = Depends(get_current_user),
                             db: AsyncSession = Depends(get_db)):
    """Record (or clear) a payment of an opening item made in cash / by bank outside a matched statement line."""
    _require_manager(user)
    ent = scope_for(user)
    row = await _item_or_404(db, ent, item_id)
    company = await company_or_404(db, ent, row.company_id)
    before = audit.snapshot(row, _ITEM_FIELDS)
    if body.paid_on is not None:
        if body.paid_on < company.books_start:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "A payment before the books start is already in the "
                                                             "opening balances; delete the item instead")
        if body.paid_account:
            chart = await ensure_chart(db, company)
            account = chart.get(body.paid_account)
            if account is None or not account.active or account.type != "asset" or account.locked:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, "Pay through an active, unlocked cash or bank account")
    row.paid_on = body.paid_on
    row.paid_via = body.paid_via if body.paid_on else ""
    row.paid_account = body.paid_account if body.paid_on else ""
    audit.record(db, label=_item_label(row), ent=ent, company_id=company.id, user=user, entity="opening_item",
                 entity_id=row.id, action="update", before=before, after=audit.snapshot(row, _ITEM_FIELDS))
    touch_books(company)
    await db.flush()
    await bank.rematch(db, company)
    await db.commit()
    return OpeningItemOut.model_validate(row)


@router.delete("/opening-items/{item_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_opening_item(item_id: uuid.UUID, user: User = Depends(get_current_user),
                              db: AsyncSession = Depends(get_db)):
    _require_manager(user)
    ent = scope_for(user)
    row = await _item_or_404(db, ent, item_id)
    company = await company_or_404(db, ent, row.company_id)
    confirmed = (await db.execute(select(AABankLine.id).where(
        AABankLine.enterprise_id == ent, AABankLine.company_id == company.id, AABankLine.match_kind == "document",
        AABankLine.matched_ref == str(row.id), AABankLine.match_status == "confirmed", bank.live_line()).limit(1)
    )).first()
    if confirmed:
        raise HTTPException(status.HTTP_409_CONFLICT, "A confirmed bank line pays this item; reject that match on the "
                                                      "Bank page first")
    row.deleted_at = datetime.now(timezone.utc)
    audit.record(db, label=_item_label(row), ent=ent, company_id=company.id, user=user, entity="opening_item",
                 entity_id=row.id, action="delete", before=audit.snapshot(row, _ITEM_FIELDS), after=None)
    touch_books(company)
    await db.flush()
    await bank.rematch(db, company)
    await db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)
