"""Write a version's workbook: one tab per form and ledger; changes since the last version in pastel yellow."""

from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.accounting.compute.specs import (
    FORM_NOTES, FORM_ORDER, FORM_TITLES, LEDGER_COLUMNS, LEDGER_ORDER, PAYABLE_ITEM, WHOLE_PESO_FORMS, form_items,
)
from app.services.spreadsheet_safe import csv_safe as _safe  # document text is untrusted: never let it become a formula

YELLOW = PatternFill("solid", fgColor="FFF4B3")
DARK = Font(color="1F1F1F")
STRUCK = Font(color="1F1F1F", strike=True)
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
RED = Font(bold=True, color="B42318")
MONEY = "#,##0.00"
WHOLE = "#,##0"
AUTHOR = "AutoAccounting"


def _number(text):
    if text is None or text == "":
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, TypeError):
        return text


def _shown(text) -> str:
    return "blank" if text in (None, "") else str(text)


def _mark(cell, before, after, parent: int, version: int) -> None:
    cell.fill = YELLOW
    cell.font = DARK
    cell.comment = Comment(f"v{parent}: {_shown(before)} → v{version}: {_shown(after)}", AUTHOR)


def _summary(wb, header, snapshot, warnings, review_open, due, version, parent_version, highlighting) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"AutoAccounting: {header['company']}"
    ws["A1"].font = TITLE
    version_text = f"v{version}" + (f", changes since v{parent_version} in yellow" if highlighting else "")
    row = 3
    for label, value in (("TIN", header["tin"]), ("Period", header["period"]), ("Version", version_text),
                         ("Generated", header["generated"])):
        ws.cell(row, 1, label).font = BOLD
        ws.cell(row, 2, _safe(value))
        row += 1
    if review_open:
        ws.cell(row, 1, f"DRAFT: {review_open} document(s) still need review").font = RED
        row += 1
    row += 1
    for col, text in enumerate(("Return", "Amount payable / (excess)", "Due date"), start=1):
        ws.cell(row, col, text).font = BOLD
    row += 1
    forms = snapshot["forms"]
    for form in FORM_ORDER:
        if form in PAYABLE_ITEM and form in forms:
            ws.cell(row, 1, FORM_TITLES[form])
            cell = ws.cell(row, 2, _number(forms[form].get(PAYABLE_ITEM[form])))
            cell.number_format = WHOLE if form in WHOLE_PESO_FORMS else MONEY
            ws.cell(row, 3, due.get(form, ""))
            row += 1
    if warnings:
        row += 1
        ws.cell(row, 1, "Warnings").font = BOLD
        row += 1
        for warning in warnings:
            ws.cell(row, 1, _safe(f"• {warning}"))
            row += 1
    for col, width in (("A", 64), ("B", 30), ("C", 14)):
        ws.column_dimensions[col].width = width


def _form_sheet(wb, form, values, changes, version, parent_version) -> None:
    ws = wb.create_sheet(form)
    ws["A1"] = FORM_TITLES[form]
    ws["A1"].font = TITLE
    for col, text in enumerate(("Item", "Description", "Amount", "How it's computed"), start=1):
        ws.cell(3, col, text).font = BOLD
    notes = FORM_NOTES.get(form, {})
    for offset, (key, label) in enumerate(form_items(form, values)):
        row = 4 + offset
        ws.cell(row, 1, key)
        ws.cell(row, 2, label)
        year = key.endswith("_YEAR")  # a year cell, not an amount
        cell = ws.cell(row, 3, values[key] if year else _number(values[key]))
        if isinstance(cell.value, Decimal):
            cell.number_format = WHOLE if form in WHOLE_PESO_FORMS else MONEY
        ws.cell(row, 4, notes.get(key, ""))
        if key in changes:
            _mark(cell, changes[key][0], changes[key][1], parent_version, version)
    ws.freeze_panes = "A4"
    for col, width in (("A", 10), ("B", 60), ("C", 18), ("D", 30)):
        ws.column_dimensions[col].width = width


def _ledger_sheet(wb, name, rows, change, version, parent_version) -> None:
    ws = wb.create_sheet(name)
    columns = LEDGER_COLUMNS[name]
    for col, (_, label, _) in enumerate(columns, start=1):
        ws.cell(1, col, label).font = BOLD
    added = set(change["added"])
    for offset, row in enumerate(rows):
        r = 2 + offset
        changed = change["changed"].get(row["key"], {})
        for col, (key, _, kind) in enumerate(columns, start=1):
            cell = ws.cell(r, col, _number(row.get(key)) if kind == "money" else _safe(row.get(key, "")))
            if kind == "money":
                cell.number_format = MONEY
            if row["key"] in added:
                cell.fill = YELLOW
                cell.font = DARK
            elif key in changed:
                _mark(cell, changed[key][0], changed[key][1], parent_version, version)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    for col, (key, _, kind) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = 16 if kind == "money" else (30 if key in ("name", "address", "counterparty") else 14)


def _change_log(wb, diff, version, parent_version) -> None:
    ws = wb.create_sheet("Change log")
    ws["A1"] = f"Changes from v{parent_version} to v{version}"
    ws["A1"].font = TITLE
    for col, text in enumerate(("Tab", "Item / row", "Field", "Before", "After"), start=1):
        ws.cell(3, col, text).font = BOLD
    row = 4
    for form, items in diff["forms"].items():
        for key, (before, after) in items.items():
            for col, value in enumerate((form, key, "value", _shown(before), _shown(after)), start=1):
                ws.cell(row, col, _safe(value))
            row += 1
    for name, change in diff["ledgers"].items():
        for key in change["added"]:
            for col, value in enumerate((name, key, "(new row)", "", ""), start=1):
                ws.cell(row, col, _safe(value)).fill = YELLOW
            row += 1
        for key, columns in change["changed"].items():
            for field, (before, after) in columns.items():
                for col, value in enumerate((name, key, field, _shown(before), _shown(after)), start=1):
                    ws.cell(row, col, _safe(value))
                row += 1
        for key in change["removed"]:
            for col, value in enumerate((name, key, "(row removed)", "", ""), start=1):
                cell = ws.cell(row, col, _safe(value))
                cell.fill = YELLOW
                cell.font = STRUCK
            row += 1
    for col, width in (("A", 14), ("B", 40), ("C", 18), ("D", 18), ("E", 18)):
        ws.column_dimensions[col].width = width


def write_workbook(path: Path, *, header: dict, snapshot: dict, warnings: list[str], review_open: int,
                   due: dict[str, str], version: int, diff: dict | None = None, parent_version: int | None = None,
                   clean: bool = False) -> None:
    highlighting = bool(diff) and not clean and parent_version is not None
    active = diff if highlighting else {"forms": {}, "ledgers": {}}
    wb = Workbook()
    _summary(wb, header, snapshot, warnings, review_open, due, version, parent_version, highlighting)
    for form in FORM_ORDER:
        if form in snapshot["forms"]:
            _form_sheet(wb, form, snapshot["forms"][form], active["forms"].get(form, {}), version, parent_version)
    for name in LEDGER_ORDER:
        if name in snapshot["ledgers"]:
            change = active["ledgers"].get(name, {"added": [], "changed": {}, "removed": []})
            _ledger_sheet(wb, name, snapshot["ledgers"][name], change, version, parent_version)
    if highlighting:
        _change_log(wb, diff, version, parent_version)
    wb.save(path)
