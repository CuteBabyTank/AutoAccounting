"""Live reports as tables, XLSX (openpyxl, ``csv_safe``) and PDF (PyMuPDF), plus the JSON encoding of report data.

Every report is first flattened to ``(heads, rows)`` where a row is ``{"cells": [...], "bold": bool}`` and a cell is
text, a Decimal (money, right-aligned) or a date.
"""

import io
import re
import uuid
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.accounting.books_pdf import MARGIN, Header
from app.services.accounting.compute.reports import _invalid, entry_ref
from app.services.accounting.subledger_report import _PAGE, _cut, _header, _Writer, peso
from app.services.spreadsheet_safe import csv_safe

_MONEY_FMT = "#,##0.00"
_MONEY_WIDTH = 82
GRID_REPORTS = ("pnl", "balance-sheet", "adjusted-tb")


def jsonable(value):
    """Report data for a JSON response: Decimals as strings with two places, dates as ISO strings."""
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, dict):
        return {k: jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [jsonable(v) for v in value]
    return value


_UNBALANCED = re.compile(r"journal entry (\S+): debits")


def engine_error_rows(message: str) -> list[dict]:
    """The books engine refused to post (e.g. an unbalanced computed entry): one Invalid Journal Transactions row."""
    found = _UNBALANCED.search(message)
    entry_id = found.group(1) if found else ""
    row = _invalid(entry_id, None, "", "engine_error", message[:300])
    if not entry_id:
        row["ref"] = {"kind": "system", "id": ""}
    else:
        row["ref"] = entry_ref(entry_id)
    return [row]


def _row(cells: list, bold: bool = False) -> dict:
    return {"cells": cells, "bold": bold}


def _blank_zero(value: Decimal) -> Decimal | str:
    return value if value else ""


def _pct(value: Decimal) -> str:
    """A % change as text; a decrease in parentheses like the money columns (a leading "-" would make csv_safe
    quote the cell in Excel)."""
    return f"{value:,.2f}%" if value >= 0 else f"({-value:,.2f}%)"


def table_for(kind: str, data: dict) -> tuple[list[str], list[dict]]:
    if kind in GRID_REPORTS:
        heads = ["Account"] + [c["label"] for c in data["columns"]]
        kinds = [c["kind"] for c in data["columns"]]
        rows = []
        for r in data["rows"]:
            cells: list = [("  " * r["level"]) + r["label"]]
            for v, k in zip(r["values"], kinds):
                cells.append("" if v is None else (_pct(v) if k == "pct" else v))
            rows.append(_row(cells, r["kind"] != "account"))
        return heads, rows
    if kind == "gl":
        heads = ["Date", "Book", "Reference", "Party", "Particulars", "Debit", "Credit", "Balance"]
        rows = []
        for a in data["accounts"]:
            rows.append(_row([f"{a['code']} {a['name']}", "", "", "", "", "", "", ""], True))
            rows.append(_row(["", "", "", "", "Opening balance", "", "", a["opening"]]))
            for l in a["lines"]:
                rows.append(_row([l["date"], l["book"], l["reference"], l["party"], l["particulars"],
                                  _blank_zero(l["debit"]), _blank_zero(l["credit"]), l["balance"]]))
            rows.append(_row(["", "", "", "", "Closing balance", "", "", a["closing"]], True))
        return heads, rows
    if kind == "transactions":
        heads = ["Date", "Book", "Reference", "Party", "Particulars", "Accounts", "Amount"]
        return heads, [_row([r["date"], r["book"], r["reference"], r["party"], r["particulars"],
                             ", ".join(r["accounts"]), r["amount"]]) for r in data["rows"]]
    if kind == "adjusting":
        heads = ["Date", "Kind", "Reference", "Particulars", "Account", "Debit", "Credit"]
        rows = []
        for e in data["rows"]:
            rows.append(_row([e["date"], e["kind"], e["reference"], e["particulars"], "", "", ""], True))
            rows += [_row(["", "", "", "", f"{l['code']} {l['name']}", _blank_zero(l["debit"]),
                           _blank_zero(l["credit"])]) for l in e["lines"]]
        return heads, rows
    if kind == "invalid":
        heads = ["Date", "Entry", "Reference", "Problem", "Detail"]
        return heads, [_row([r["date"] or "", r["entry_id"], r["reference"], r["label"], r["detail"]])
                       for r in data["rows"]]
    raise ValueError(f"Unknown report {kind!r}")


def report_xlsx(header: Header, title: str, subtitle: str, heads: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = csv_safe(title[:31].replace("/", "-")) or "Report"
    bold = Font(bold=True)
    ws["A1"] = csv_safe(header.name)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = csv_safe(title)
    ws["A2"].font = bold
    ws["A3"] = csv_safe(subtitle)
    for col, text in enumerate(heads, start=1):
        ws.cell(5, col, text).font = bold
    for n, row in enumerate(rows, start=6):
        for col, value in enumerate(row["cells"], start=1):
            cell = ws.cell(n, col, csv_safe(value) if isinstance(value, str) else value)
            if isinstance(value, Decimal):
                cell.number_format = _MONEY_FMT
            if isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
            if row["bold"]:
                cell.font = bold
    ws.column_dimensions["A"].width = 44
    for letter in "BCDEFGHIJKLMNOPQRSTUVWXYZ":
        ws.column_dimensions[letter].width = 16
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(value) -> str:
    if isinstance(value, Decimal):
        return peso(value)
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


_PDF_SAFE = str.maketrans({"₱": "PHP ", "–": "-", "—": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "·": "-"})


def pdf_text(value) -> str:
    """Cell text for the base-14 Helvetica font (no peso sign or typographic dashes)."""
    return _text(value).translate(_PDF_SAFE)


def layout(heads: list[str], rows: list[dict], width: float) -> list[tuple[float, float, str]]:
    """(x, width, align) per column: money columns fixed and right-aligned, text columns share the rest."""
    money = [any(isinstance(r["cells"][i], Decimal) for r in rows if i < len(r["cells"])) or h.endswith("change")
             for i, h in enumerate(heads)]
    text_cols = [i for i, m in enumerate(money) if not m]
    weights = {i: (3 if i == 0 else 1) for i in text_cols}
    free = width - _MONEY_WIDTH * sum(money)
    unit = free / max(sum(weights.values()), 1)
    out, x = [], MARGIN
    for i, is_money in enumerate(money):
        w = _MONEY_WIDTH if is_money else unit * weights[i]
        out.append((x + w - 4 if is_money else x, w - 6, "r" if is_money else "l"))
        x += w
    return out


def report_pdf(header: Header, title: str, subtitle: str, heads: list[str], rows: list[dict],
               writer: _Writer | None = None) -> bytes:
    w = writer or _Writer(_PAGE)
    _header(w, Header(pdf_text(header.name), header.tin, pdf_text(header.address), header.rdo), pdf_text(title.upper()),
            pdf_text(subtitle))
    cols = layout(heads, rows, _PAGE.width - 2 * MARGIN)

    def line(cells: list, bold: bool) -> None:
        w.row([(x, _cut(pdf_text(c), width) if align == "l" else pdf_text(c), align)
               for (x, width, align), c in zip(cols, cells)], bold=bold)

    line(heads, True)
    for r in rows:
        if w.y > w.rect.height - MARGIN - 13:
            w.new_page()
            line(heads, True)
        line(r["cells"], r["bold"])
    return w.bytes()
