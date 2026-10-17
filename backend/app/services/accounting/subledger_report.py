"""A/R and A/P aging (XLSX + PDF) and the customer statement PDF (PyMuPDF / openpyxl).

Names, numbers and TINs come from documents, so every spreadsheet text cell goes through ``csv_safe``.
"""

import io
from datetime import date
from decimal import Decimal

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.accounting.books_pdf import BOLD, FONT, MARGIN, Header, _fit
from app.services.accounting.compute.subledger import BUCKET_LABELS, BUCKETS, AgingRow, OpenDoc, OpenItem
from app.services.spreadsheet_safe import csv_safe

_PAGE = fitz.Rect(0, 0, 936, 612)        # 13" × 8.5" landscape: six money columns fit
_PORTRAIT = fitz.Rect(0, 0, 612, 936)
_ROW = 13
_MONEY = "#,##0.00"
TITLES = {"ar": "ACCOUNTS RECEIVABLE AGING", "ap": "ACCOUNTS PAYABLE AGING"}
PARTY = {"ar": "Customer", "ap": "Supplier"}


def _cut(text: str, width: float) -> str:
    """Shorten a table cell to ``width`` at the size rows are drawn (8 pt), so it never runs into the next column."""
    return _fit(text, width, FONT, 8)


def peso(value: Decimal) -> str:
    return f"{value:,.2f}" if value >= 0 else f"({-value:,.2f})"   # Helvetica has no peso sign


class _Writer:
    """Rows of text with right-aligned money columns, flowing onto new pages."""

    def __init__(self, rect: fitz.Rect):
        self.doc = fitz.open()
        self.rect = rect
        self.page = None
        self.y = 0.0
        self.new_page()

    def new_page(self) -> None:
        self.page = self.doc.new_page(width=self.rect.width, height=self.rect.height)
        self.y = MARGIN + 6

    def row(self, cells: list[tuple[float, str, str]], bold: bool = False, size: float = 8) -> None:
        """``cells``: (x, text, align) with align "l" (x = left edge) or "r" (x = right edge)."""
        if self.y > self.rect.height - MARGIN - _ROW:
            self.new_page()
        font = BOLD if bold else FONT
        for x, text, align in cells:
            if align == "r":
                width = fitz.get_text_length(text, fontname=font, fontsize=size)
                self.page.insert_text((x - width, self.y), text, fontname=font, fontsize=size)
            else:
                self.page.insert_text((x, self.y), text, fontname=font, fontsize=size)
        self.y += _ROW

    def text(self, value: str, bold: bool = False, size: float = 9) -> None:
        self.row([(MARGIN, _fit(value, self.rect.width - 2 * MARGIN, BOLD if bold else FONT, size), "l")], bold, size)

    def gap(self, points: float = 6) -> None:
        self.y += points

    def bytes(self) -> bytes:
        return self.doc.tobytes(garbage=3, deflate=True)


def _header(w: _Writer, header: Header, title: str, subtitle: str) -> None:
    w.text(header.name.upper(), bold=True, size=11)
    w.text("   ·   ".join(p for p in (f"TIN {header.tin}", header.address) if p.strip()), size=8)
    w.gap()
    w.text(title, bold=True, size=11)
    w.text(subtitle, size=9)
    w.gap(8)


def aging_pdf(header: Header, kind: str, as_of: date, rows: list[AgingRow], totals: dict[str, Decimal],
              items: list[OpenItem]) -> bytes:
    """Summary by contact (buckets), then the open items; A/P marks bills with no payment seen yet."""
    w = _Writer(_PAGE)
    _header(w, header, TITLES[kind], f"As of {as_of:%B %d, %Y}")
    right = _PAGE.width - MARGIN
    cols = [right - 95 * (len(BUCKETS) - i) for i in range(len(BUCKETS))] + [right]
    name_width = cols[0] - 95 - MARGIN - 70

    def heads() -> None:
        w.row([(MARGIN, PARTY[kind], "l"), (MARGIN + name_width + 8, "TIN", "l")]
              + [(x, BUCKET_LABELS[b], "r") for x, b in zip(cols, BUCKETS)] + [(cols[-1], "Total", "r")], bold=True)

    heads()
    for r in rows:
        name = r.name + (f"  ({r.unverified} unverified)" if kind == "ap" and r.unverified else "")
        w.row([(MARGIN, _cut(name, name_width), "l"), (MARGIN + name_width + 8, r.tin, "l")]
              + [(x, peso(r.buckets[b]), "r") for x, b in zip(cols, BUCKETS)] + [(cols[-1], peso(r.total), "r")])
    w.row([(MARGIN, "Total", "l")] + [(x, peso(totals[b]), "r") for x, b in zip(cols, BUCKETS)]
          + [(cols[-1], peso(totals["total"]), "r")], bold=True)
    w.gap(12)
    w.text("Open items", bold=True)
    x_no, x_date, x_due = MARGIN + 230, MARGIN + 330, MARGIN + 400
    w.row([(MARGIN, PARTY[kind], "l"), (x_no, "Document", "l"), (x_date, "Date", "l"), (x_due, "Due", "l"),
           (right - 380, "Days past due", "r"), (right - 285, "Total", "r"), (right - 190, "Withheld", "r"),
           (right - 95, "Paid", "r"), (right, "Open", "r")], bold=True)
    for i in items:
        d = i.doc
        note = "  * unverified" if kind == "ap" and d.payment_unverified else ""
        w.row([(MARGIN, _cut(d.contact_name, 220), "l"), (x_no, _cut(d.doc_number + note, 95), "l"),
               (x_date, d.doc_date.isoformat(), "l"), (x_due, i.due.isoformat(), "l"),
               (right - 380, str(max(i.days_past_due, 0)), "r"), (right - 285, peso(d.total), "r"),
               (right - 190, peso(d.withheld), "r"), (right - 95, peso(i.paid), "r"), (right, peso(i.open), "r")])
    if kind == "ap" and any(i.doc.payment_unverified for i in items):
        w.gap()
        w.text("* unverified: a bill on credit terms with no receipt or bank line paying it yet (payment_unverified).",
               size=7.5)
    if kind == "ar":
        w.gap()
        w.text("Balances are net of tax the customer withheld; that tax is tracked as 2307s to receive.", size=7.5)
    return w.bytes()


def statement_pdf(header: Header, contact: dict, as_of: date, items: list[OpenItem], cwt_pending: Decimal) -> bytes:
    """Open-item customer statement with an aging footer (``contact``: name, tin, address)."""
    w = _Writer(_PORTRAIT)
    _header(w, header, "STATEMENT OF ACCOUNT", f"As of {as_of:%B %d, %Y}")
    w.text(f"To: {contact.get('name', '')}", bold=True)
    if contact.get("tin"):
        w.text(f"TIN {contact['tin']}", size=8)
    if contact.get("address"):
        w.text(contact["address"], size=8)
    w.gap(8)
    right = _PORTRAIT.width - MARGIN
    w.row([(MARGIN, "Invoice", "l"), (MARGIN + 110, "Date", "l"), (MARGIN + 180, "Due", "l"),
           (right - 270, "Amount", "r"), (right - 180, "Tax withheld", "r"), (right - 90, "Paid", "r"),
           (right, "Balance due", "r")], bold=True)
    for i in items:
        d = i.doc
        w.row([(MARGIN, _cut(d.doc_number, 100), "l"), (MARGIN + 110, d.doc_date.isoformat(), "l"),
               (MARGIN + 180, i.due.isoformat(), "l"), (right - 270, peso(d.total), "r"),
               (right - 180, peso(d.withheld), "r"), (right - 90, peso(i.paid), "r"), (right, peso(i.open), "r")])
    total = sum((i.open for i in items), Decimal("0.00"))
    w.row([(MARGIN, "Total balance due", "l"), (right, peso(total), "r")], bold=True)
    w.gap(10)
    w.text("Aging", bold=True)
    step = (right - MARGIN) / len(BUCKETS)
    w.row([(MARGIN + step * (n + 1), BUCKET_LABELS[b], "r") for n, b in enumerate(BUCKETS)], bold=True)
    w.row([(MARGIN + step * (n + 1), peso(sum((i.open for i in items if i.bucket == b), Decimal("0.00"))), "r")
           for n, b in enumerate(BUCKETS)])
    if cwt_pending > 0:
        w.gap(8)
        w.text(f"Please send BIR Form 2307 for the tax you withheld: {peso(cwt_pending)}.", size=8)
    return w.bytes()


def aging_xlsx(header: Header, kind: str, as_of: date, rows: list[AgingRow], totals: dict[str, Decimal],
               items: list[OpenItem]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aging"
    bold = Font(bold=True)
    ws["A1"] = csv_safe(header.name)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{TITLES[kind].title()} as of {as_of.isoformat()}"
    heads = [PARTY[kind], "TIN"] + [BUCKET_LABELS[b] for b in BUCKETS] + ["Total", "Open items"]
    if kind == "ap":
        heads.append("Unverified")
    for col, text in enumerate(heads, start=1):
        ws.cell(4, col, text).font = bold
    r = 5
    for row in rows:
        values = [csv_safe(row.name), csv_safe(row.tin)] + [row.buckets[b] for b in BUCKETS] + [row.total, row.count]
        if kind == "ap":
            values.append(row.unverified)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(r, col, value)
            if isinstance(value, Decimal):
                cell.number_format = _MONEY
        r += 1
    ws.cell(r, 1, "Total").font = bold
    for col, b in enumerate(list(BUCKETS) + ["total"], start=3):
        cell = ws.cell(r, col, totals[b])
        cell.number_format, cell.font = _MONEY, bold
    for col, width in zip("ABCDEFGHIJ", (40, 14, 15, 15, 15, 15, 15, 16, 11, 11)):
        ws.column_dimensions[col].width = width

    detail = wb.create_sheet("Open items")
    heads = [PARTY[kind], "TIN", "Document", "Date", "Due", "Days past due", "Bucket", "Total", "Withheld", "Paid",
             "Open"] + (["Payment unverified"] if kind == "ap" else [])
    for col, text in enumerate(heads, start=1):
        detail.cell(1, col, text).font = bold
    for n, i in enumerate(items, start=2):
        d = i.doc
        values = [csv_safe(d.contact_name), csv_safe(d.contact_tin), csv_safe(d.doc_number), d.doc_date.isoformat(),
                  i.due.isoformat(), max(i.days_past_due, 0), BUCKET_LABELS[i.bucket], d.total, d.withheld, i.paid,
                  i.open] + (["yes" if d.payment_unverified else ""] if kind == "ap" else [])
        for col, value in enumerate(values, start=1):
            cell = detail.cell(n, col, value)
            if isinstance(value, Decimal):
                cell.number_format = _MONEY
    detail.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def aging_ledger_rows(items: list[OpenItem]) -> list[dict]:
    """A version's aging tab: one row per open item, its open amount in its bucket's column (text amounts)."""
    out = []
    for i in items:
        d = i.doc
        row = {"key": d.id, "party": d.contact_name, "tin": d.contact_tin, "contact": d.contact_key,
               "doc_number": d.doc_number, "date": d.doc_date.isoformat(), "due": i.due.isoformat(),
               "days": str(max(i.days_past_due, 0)), "total_amount": str(d.total), "withheld": str(d.withheld),
               "paid": str(i.paid), "open": str(i.open), "unverified": "yes" if d.payment_unverified else ""}
        row.update({b: str(i.open) if b == i.bucket else "0.00" for b in BUCKETS})
        out.append(row)
    return out


def items_from_ledger(kind: str, rows: list[dict]) -> list[OpenItem]:
    """The open items a version's aging tab was written from (to print its PDF later, unchanged)."""
    out = []
    for r in rows:
        due, on = date.fromisoformat(r["due"]), date.fromisoformat(r["date"])
        bucket = next((b for b in BUCKETS if Decimal(r.get(b) or "0") != 0), "current")
        doc = OpenDoc(id=r["key"], kind=kind, contact_key=r.get("contact") or r["party"], contact_name=r["party"],
                      contact_tin=r.get("tin", ""), doc_number=r.get("doc_number", ""), doc_date=on,
                      total=Decimal(r["total_amount"]), withheld=Decimal(r["withheld"]),
                      payment_unverified=r.get("unverified") == "yes")
        out.append(OpenItem(doc, due, Decimal(r["paid"]), Decimal(r["open"]), int(r.get("days") or 0), bucket))
    return out
