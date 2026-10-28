"""Form filler service: parse compact Excel → validate → render merged PDF.

Two workbook perspectives are supported (auto-detected on upload):

* **Payee Form**: payee identity once at the top, then one row per *payor*.
* **Payor Form**: payor identity once at the top, then one row per *payee*.

Both use the same compact, single-sheet layout::

    Row 1  <Title banner: "Payee Form" / "Payor Form">
    Row 2  <identity header: tin | name | registered_address | zip | [foreign] | from_date | to_date>
    Row 3  <identity data>
    Row 4  (blank)
    Row 5  (decorative month labels above the amount columns)
    Row 6  <entity header: name | registered_address | zip | tin | block | description | atc | m1 | m2 | m3 | total | tax_withheld>
    Row 7+ <one row per entity; a row with a BLANK name but amount data is a
            continuation line item attached to the entity above it>

A single period (``from_date`` / ``to_date``) at the top applies to every
generated 2307.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from app.schemas.form_filler import FilledForm, LineItem, Payee, Payor

PAYEE_KIND = "payee"
PAYOR_KIND = "payor"

# Human-facing sheet label used in validation error rows.
SHEET_LABELS = {PAYEE_KIND: "PayeeForm", PAYOR_KIND: "PayorForm"}

LINE_ITEM_HEADERS = [
    "block",
    "description",
    "atc",
    "m1",
    "m2",
    "m3",
    "total",
    "tax_withheld",
]

EWT_BLOCK = "ewt"
MP_BLOCK = "mp"
MAX_LINE_ITEMS_PER_BLOCK = 10

# The constant-identity block sits on rows 2 (header) + 3 (data) of every sheet.
IDENTITY_DATA_ROW = 3

# Canonical header tokens → accepted aliases (all matched case-insensitively
# after whitespace/parenthesis normalization). Lets both the generated template
# (canonical tokens) and prettier hand-made sheets (human labels) round-trip.
_HEADER_ALIASES = {
    "payee_tin": "payee_tin", "payee tin": "payee_tin",
    "payee_name": "payee_name", "payee name": "payee_name",
    "payee_registered_address": "payee_registered_address",
    "payee registered address": "payee_registered_address",
    "payee_zip_code": "payee_zip_code", "payee zip code": "payee_zip_code",
    "payee zip": "payee_zip_code",
    "payee_foreign_address": "payee_foreign_address",
    "payee foreign address": "payee_foreign_address",
    "payor_tin": "payor_tin", "payor tin": "payor_tin",
    "payor_name": "payor_name", "payor name": "payor_name",
    "payor_registered_address": "payor_registered_address",
    "payor registered address": "payor_registered_address",
    "payor_zip_code": "payor_zip_code", "payor zip code": "payor_zip_code",
    "payor zip": "payor_zip_code",
    "from_date": "from_date", "from date": "from_date",
    "to_date": "to_date", "to date": "to_date",
    "block": "block",
    "description": "description",
    "atc": "atc",
    "m1": "m1", "january": "m1",
    "m2": "m2", "february": "m2",
    "m3": "m3", "march": "m3",
    "total": "total", "qtr": "total", "quarter": "total",
    "tax_withheld": "tax_withheld", "tax withheld": "tax_withheld",
    "withheld": "tax_withheld",
}


@dataclass
class ParseResult:
    kind: str  # PAYEE_KIND | PAYOR_KIND
    period_from: str
    period_to: str
    payee: Payee | None  # the constant identity when kind == PAYEE_KIND
    payor: Payor | None  # the constant identity when kind == PAYOR_KIND
    forms: list[FilledForm] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def _load_workbook(file_bytes: bytes) -> Workbook:
    try:
        return load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    except Exception as e:
        raise ValueError(f"Could not read workbook: {e}") from e


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(value: object) -> str:
    """Normalize a header cell to its canonical token."""
    s = _str(value).lower().strip().strip("()").strip()
    s = re.sub(r"\s+", " ", s)
    return _HEADER_ALIASES.get(s, s.replace(" ", "_"))


def _row_dict(headers: list[str], row: tuple) -> dict[str, object]:
    return {h: v for h, v in zip(headers, row)}


def _norm_set(row: tuple) -> set[str]:
    return {_norm_header(c) for c in row if not _is_blank(c)}


def _next_non_blank(rows: list[tuple], start: int) -> tuple[int, tuple] | None:
    for idx in range(start, len(rows)):
        if not all(_is_blank(c) for c in rows[idx]):
            return idx, rows[idx]
    return None


def _build_line_item(row: dict[str, object], excel_row: int) -> LineItem | None:
    desc = row.get("description")
    atc = row.get("atc")
    tw = row.get("tax_withheld")
    m1_raw = row.get("m1")
    m2_raw = row.get("m2")
    m3_raw = row.get("m3")
    total_raw = row.get("total")
    if all(_is_blank(v) for v in (desc, atc, tw, m1_raw, m2_raw, m3_raw, total_raw)):
        return None
    try:
        m1f = float(m1_raw) if not _is_blank(m1_raw) else 0.0
        m2f = float(m2_raw) if not _is_blank(m2_raw) else 0.0
        m3f = float(m3_raw) if not _is_blank(m3_raw) else 0.0
        totalf = float(total_raw) if not _is_blank(total_raw) else m1f + m2f + m3f
        twf = float(tw) if not _is_blank(tw) else 0.0
    except (TypeError, ValueError):
        m1f = m2f = m3f = totalf = twf = 0.0
    return LineItem(
        description=_str(desc),
        atc=_str(atc),
        m1=m1f, m2=m2f, m3=m3f,
        total=totalf,
        tax_withheld=twf,
        excel_row=excel_row,
    )


# --------------------------------------------------------------------------- #
# Detection + parsing
# --------------------------------------------------------------------------- #
def _detect_kind(ws, rows: list[tuple]) -> str:
    """Decide whether the sheet is a Payee Form or a Payor Form."""
    title = _str(ws.cell(row=1, column=1).value).lower()
    if "payee form" in title:
        return PAYEE_KIND
    if "payor form" in title:
        return PAYOR_KIND
    # Fallback: the repeating-entity header reveals the perspective.
    for raw in rows:
        norm = _norm_set(raw)
        if "block" in norm and "description" in norm:
            if "payor_name" in norm:
                return PAYEE_KIND
            if "payee_name" in norm:
                return PAYOR_KIND
    raise ValueError(
        "Could not tell the form type. Expected a 'Payee Form' or 'Payor Form' workbook"
    )


def _parse_compact(rows: list[tuple], kind: str) -> ParseResult:
    const_prefix = "payee" if kind == PAYEE_KIND else "payor"
    var_prefix = "payor" if kind == PAYEE_KIND else "payee"

    # --- Constant identity block (header + data row) ---
    id_hdr_idx: int | None = None
    id_headers: list[str] = []
    for idx, raw in enumerate(rows):
        norm = _norm_set(raw)
        if (f"{const_prefix}_tin" in norm and f"{const_prefix}_name" in norm
                and "block" not in norm):
            id_hdr_idx = idx
            id_headers = [_norm_header(c) for c in raw]
            break
    if id_hdr_idx is None:
        raise ValueError(f"Sheet has no {const_prefix} identity section")

    nxt = _next_non_blank(rows, id_hdr_idx + 1)
    if nxt is None:
        raise ValueError(f"{const_prefix.capitalize()} section has no data row")
    id_data_idx, id_data_row = nxt
    id_dict = _row_dict(id_headers, id_data_row)
    period_from = _str(id_dict.get("from_date"))
    period_to = _str(id_dict.get("to_date"))

    constant_payee: Payee | None = None
    constant_payor: Payor | None = None
    if kind == PAYEE_KIND:
        constant_payee = Payee(
            tin=_str(id_dict.get("payee_tin")),
            name=_str(id_dict.get("payee_name")),
            registered_address=_str(id_dict.get("payee_registered_address")),
            zip_code=_str(id_dict.get("payee_zip_code")),
            foreign_address=_str(id_dict.get("payee_foreign_address")),
        )
    else:
        constant_payor = Payor(
            tin=_str(id_dict.get("payor_tin")),
            name=_str(id_dict.get("payor_name")),
            registered_address=_str(id_dict.get("payor_registered_address")),
            zip_code=_str(id_dict.get("payor_zip_code")),
        )

    # --- Repeating entity table ---
    ent_hdr_idx: int | None = None
    ent_headers: list[str] = []
    for idx in range(id_data_idx + 1, len(rows)):
        norm = _norm_set(rows[idx])
        if "block" in norm and "description" in norm and f"{var_prefix}_name" in norm:
            ent_hdr_idx = idx
            ent_headers = [_norm_header(c) for c in rows[idx]]
            break
    if ent_hdr_idx is None:
        raise ValueError(f"Sheet has no {var_prefix} table (entity rows)")

    name_key = f"{var_prefix}_name"

    @dataclass
    class _Block:
        name: str
        registered_address: str
        zip_code: str
        tin: str
        excel_row: int
        ewt: list[LineItem] = field(default_factory=list)
        mp: list[LineItem] = field(default_factory=list)

    blocks: list[_Block] = []
    current: _Block | None = None
    for idx in range(ent_hdr_idx + 1, len(rows)):
        raw = rows[idx]
        if all(_is_blank(c) for c in raw):
            continue
        norm = _norm_set(raw)
        # Skip a stray repeated header row.
        if "block" in norm and "description" in norm and name_key in norm:
            continue
        rowd = _row_dict(ent_headers, raw)
        name = _str(rowd.get(name_key))
        if name:
            current = _Block(
                name=name,
                registered_address=_str(rowd.get(f"{var_prefix}_registered_address")),
                zip_code=_str(rowd.get(f"{var_prefix}_zip_code")),
                tin=_str(rowd.get(f"{var_prefix}_tin")),
                excel_row=idx + 1,
            )
            blocks.append(current)
        li = _build_line_item(rowd, idx + 1)
        if li is not None:
            if current is None:
                raise ValueError(
                    f"Line item on row {idx + 1} has no {var_prefix} above it. "
                    f"the first {var_prefix} row must include a name"
                )
            block_kind = _str(rowd.get("block")).lower()
            (current.mp if block_kind == MP_BLOCK else current.ewt).append(li)

    forms: list[FilledForm] = []
    for b in blocks:
        if kind == PAYEE_KIND:
            assert constant_payee is not None
            payee = constant_payee
            payor = Payor(tin=b.tin, name=b.name,
                          registered_address=b.registered_address, zip_code=b.zip_code)
        else:
            assert constant_payor is not None
            payor = constant_payor
            payee = Payee(tin=b.tin, name=b.name,
                          registered_address=b.registered_address,
                          zip_code=b.zip_code, foreign_address="")
        forms.append(FilledForm(
            excel_row=b.excel_row, payee=payee, payor=payor,
            period_from=period_from, period_to=period_to,
            ewt=b.ewt, money_payments=b.mp,
        ))

    return ParseResult(
        kind=kind, period_from=period_from, period_to=period_to,
        payee=constant_payee, payor=constant_payor, forms=forms,
    )


def parse_workbook(file_bytes: bytes) -> ParseResult:
    """Detect the perspective and parse the compact workbook."""
    wb = _load_workbook(file_bytes)
    if not wb.sheetnames:
        raise ValueError("Workbook has no sheets")
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Worksheet is empty")
    kind = _detect_kind(ws, rows)
    result = _parse_compact(rows, kind)
    if not result.forms:
        raise ValueError("Sheet has no entity rows to generate forms from")
    return result


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
from app.schemas.form_filler import FieldError
from app.services.atc_codes import lookup_atc

_TIN_RE = re.compile(r"^\d{3}-\d{3}-\d{3}-\d{3,5}$")
_DATE_RE = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")
_ZIP_RE = re.compile(r"^\d{4}$")
_TOTAL_TOLERANCE = 0.01


def _err(sheet: str, row: int, column: str, code: str, message: str) -> FieldError:
    return FieldError(sheet=sheet, row=row, column=column, code=code, message=message)


def _check_tin(value: str, sheet: str, row: int, column: str,
               errors: list[FieldError]) -> None:
    if not value:
        errors.append(_err(sheet, row, column, "REQUIRED", "Field is required"))
        return
    if not _TIN_RE.match(value):
        errors.append(_err(sheet, row, column, "TIN_FORMAT",
                           "TIN must be formatted XXX-XXX-XXX-XXX (last group may be 3-5 digits)"))


def _check_zip(value: str, sheet: str, row: int, column: str,
               errors: list[FieldError]) -> None:
    if not value:
        errors.append(_err(sheet, row, column, "REQUIRED", "Field is required"))
        return
    if not _ZIP_RE.match(value):
        errors.append(_err(sheet, row, column, "ZIP_FORMAT", "ZIP code must be 4 digits"))


def _check_required(value: str, sheet: str, row: int, column: str,
                    errors: list[FieldError]) -> None:
    if not value:
        errors.append(_err(sheet, row, column, "REQUIRED", "Field is required"))


def _check_date(value: str, sheet: str, row: int, column: str,
                errors: list[FieldError]) -> tuple[int, int, int] | None:
    if not value:
        errors.append(_err(sheet, row, column, "REQUIRED", "Field is required"))
        return None
    m = _DATE_RE.match(value)
    if not m:
        errors.append(_err(sheet, row, column, "DATE_FORMAT", "Date must be MM/DD/YYYY"))
        return None
    mm, dd, yyyy = (int(g) for g in m.groups())
    if not (1 <= mm <= 12 and 1 <= dd <= 31):
        errors.append(_err(sheet, row, column, "DATE_FORMAT", "Date must be MM/DD/YYYY"))
        return None
    return (yyyy, mm, dd)


def _check_line_item(li: LineItem, sheet: str, fallback_row: int,
                     errors: list[FieldError]) -> None:
    row = li.excel_row or fallback_row
    if not li.description or not li.atc:
        errors.append(_err(sheet, row, "description", "LINE_ITEM_PARTIAL",
                           "Incomplete line item. Fill in every required cell or leave the whole row blank."))
    else:
        if lookup_atc(li.atc) is None:
            errors.append(_err(sheet, row, "atc", "ATC_UNKNOWN",
                               f"Unknown ATC code: {li.atc}"))
    for f in ("m1", "m2", "m3", "total", "tax_withheld"):
        if getattr(li, f) < 0:
            errors.append(_err(sheet, row, f, "AMOUNT_NEGATIVE", "Amount cannot be negative"))
    sum_m = li.m1 + li.m2 + li.m3
    if abs(li.total - sum_m) > _TOTAL_TOLERANCE:
        errors.append(_err(sheet, row, "total", "TOTAL_MISMATCH",
                           "Total does not match m1+m2+m3"))


def validate(result: ParseResult) -> list[FieldError]:
    errors: list[FieldError] = []
    sheet = SHEET_LABELS[result.kind]
    var_prefix = "payor" if result.kind == PAYEE_KIND else "payee"

    # Single shared period (top of sheet).
    d_from = _check_date(result.period_from, sheet, IDENTITY_DATA_ROW, "from_date", errors)
    d_to = _check_date(result.period_to, sheet, IDENTITY_DATA_ROW, "to_date", errors)
    if d_from and d_to and d_to < d_from:
        errors.append(_err(sheet, IDENTITY_DATA_ROW, "to_date", "DATE_RANGE",
                           "Period 'To' is before 'From'"))

    # Constant identity (filled once, row 3).
    if result.kind == PAYEE_KIND:
        p = result.payee
        assert p is not None
        _check_tin(p.tin, sheet, IDENTITY_DATA_ROW, "payee_tin", errors)
        _check_required(p.name, sheet, IDENTITY_DATA_ROW, "payee_name", errors)
        _check_required(p.registered_address, sheet, IDENTITY_DATA_ROW,
                        "payee_registered_address", errors)
        _check_zip(p.zip_code, sheet, IDENTITY_DATA_ROW, "payee_zip_code", errors)
    else:
        p = result.payor
        assert p is not None
        _check_tin(p.tin, sheet, IDENTITY_DATA_ROW, "payor_tin", errors)
        _check_required(p.name, sheet, IDENTITY_DATA_ROW, "payor_name", errors)
        _check_required(p.registered_address, sheet, IDENTITY_DATA_ROW,
                        "payor_registered_address", errors)
        _check_zip(p.zip_code, sheet, IDENTITY_DATA_ROW, "payor_zip_code", errors)

    if not result.forms:
        errors.append(_err(sheet, IDENTITY_DATA_ROW, "*", "EMPTY_FORMS",
                           "No forms to generate"))
        return errors

    for form in result.forms:
        r = form.excel_row
        ent = form.payor if result.kind == PAYEE_KIND else form.payee
        _check_tin(ent.tin, sheet, r, f"{var_prefix}_tin", errors)
        _check_required(ent.name, sheet, r, f"{var_prefix}_name", errors)
        _check_required(ent.registered_address, sheet, r,
                        f"{var_prefix}_registered_address", errors)
        _check_zip(ent.zip_code, sheet, r, f"{var_prefix}_zip_code", errors)
        if not form.ewt and not form.money_payments:
            errors.append(_err(sheet, r, "*", "NO_LINE_ITEMS",
                               "Form must have at least one line item"))
        if len(form.ewt) > MAX_LINE_ITEMS_PER_BLOCK:
            errors.append(_err(sheet, r, "block", "TOO_MANY_LINE_ITEMS",
                               f"At most {MAX_LINE_ITEMS_PER_BLOCK} ewt line items per form"))
        if len(form.money_payments) > MAX_LINE_ITEMS_PER_BLOCK:
            errors.append(_err(sheet, r, "block", "TOO_MANY_LINE_ITEMS",
                               f"At most {MAX_LINE_ITEMS_PER_BLOCK} mp line items per form"))
        for li in form.ewt:
            _check_line_item(li, sheet, r, errors)
        for li in form.money_payments:
            _check_line_item(li, sheet, r, errors)
    return errors


# --------------------------------------------------------------------------- #
# Template builder
# --------------------------------------------------------------------------- #
def build_template(kind: str = PAYEE_KIND) -> bytes:
    """Build a prettied, single-sheet compact template for the given perspective."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if kind not in (PAYEE_KIND, PAYOR_KIND):
        raise ValueError(f"Unknown template kind: {kind}")

    const_prefix = "payee" if kind == PAYEE_KIND else "payor"
    var_prefix = "payor" if kind == PAYEE_KIND else "payee"
    title = "Payee Form" if kind == PAYEE_KIND else "Payor Form"

    if kind == PAYEE_KIND:
        identity_headers = ["payee_tin", "payee_name", "payee_registered_address",
                            "payee_zip_code", "payee_foreign_address", "from_date", "to_date"]
        identity_data = ["123-456-789-000", "Your Company Inc.",
                         "123 Sample St., Makati City", "1226", "", "01/01/2026", "03/31/2026"]
    else:
        identity_headers = ["payor_tin", "payor_name", "payor_registered_address",
                            "payor_zip_code", "from_date", "to_date"]
        identity_data = ["999-888-777-000", "Your Company Inc.",
                         "1 Main St., Quezon City", "1100", "01/01/2026", "03/31/2026"]

    entity_headers = [f"{var_prefix}_name", f"{var_prefix}_registered_address",
                      f"{var_prefix}_zip_code", f"{var_prefix}_tin",
                      "block", "description", "atc", "m1", "m2", "m3", "total", "tax_withheld"]
    ncols = len(entity_headers)  # 12

    # Sample rows demonstrate the continuation rule (blank name = same entity).
    sample_rows = [
        ["Sample Corp", "1 Main St., Quezon City", "1100", "111-222-333-000",
         "ewt", "Professional fees", "WI010", 10000, 10000, 10000, 30000, 1500],
        ["", "", "", "",
         "ewt", "Consulting (same payor)", "WI010", 5000, 0, 0, 5000, 250],
        ["Second Entity Inc.", "2 Rizal Ave., Pasig", "1600", "444-555-666-000",
         "mp", "Supplies", "WC158", 2000, 0, 0, 2000, 20],
    ]
    BLANK_ENTITY_ROWS = 8

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_LABELS[kind]

    col_widths = [26, 30, 14, 20, 8, 26, 10, 12, 12, 12, 14, 12]
    for idx, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    thin = Side(style="thin", color="FF9CA3AF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_font = Font(bold=True, size=14, color="FF1F2937")
    group_font = Font(bold=True, size=10, color="FF374151")
    header_font = Font(bold=True, size=10, color="FF111827")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    title_fill = PatternFill("solid", fgColor="FFFDE68A")
    identity_fill = PatternFill("solid", fgColor="FFBFDBFE")
    entity_fill = PatternFill("solid", fgColor="FFBBF7D0")
    month_fill = PatternFill("solid", fgColor="FFFBCFE8")

    def _style_row(row: int, fill, font, align, n: int) -> None:
        for c in range(1, n + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = fill
            cell.font = font
            cell.alignment = align
            cell.border = border

    # Row 1: title banner.
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    _style_row(1, title_fill, title_font, center, ncols)
    ws.row_dimensions[1].height = 28

    # Row 2 / 3: identity header + data.
    for c, h in enumerate(identity_headers, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = header_font
        cell.alignment = left
        cell.border = border
        cell.fill = identity_fill
    for c, v in enumerate(identity_data, start=1):
        cell = ws.cell(row=3, column=c, value=v)
        cell.alignment = left
        cell.border = border
    ws.row_dimensions[2].height = 20

    # Row 5: decorative month labels above the amount columns (m1..tax_withheld).
    month_labels = {8: "JANUARY", 9: "FEBRUARY", 10: "MARCH",
                    11: "QTR (total)", 12: "WITHHELD"}
    for c, label in month_labels.items():
        cell = ws.cell(row=5, column=c, value=label)
        cell.font = group_font
        cell.alignment = center
        cell.fill = month_fill
        cell.border = border

    # Row 6: entity (repeating) header.
    for c, h in enumerate(entity_headers, start=1):
        cell = ws.cell(row=6, column=c, value=h)
        cell.font = header_font
        cell.alignment = left
        cell.border = border
        cell.fill = entity_fill
    ws.row_dimensions[6].height = 20

    # Rows 7+: sample data then blank rows.
    next_row = 7
    for r in sample_rows:
        for c, v in enumerate(r, start=1):
            cell = ws.cell(row=next_row, column=c, value=v)
            cell.alignment = left
            cell.border = border
        next_row += 1
    for _ in range(BLANK_ENTITY_ROWS):
        for c in range(1, ncols + 1):
            ws.cell(row=next_row, column=c).border = border
        next_row += 1

    ws.freeze_panes = "A7"  # keep identity + headers visible while scrolling

    # Comments.
    def _col(headers: list[str], name: str) -> int:
        return headers.index(name) + 1

    tin_letter = get_column_letter(_col(identity_headers, f"{const_prefix}_tin"))
    ws[f"{tin_letter}2"].comment = Comment(
        "Format XXX-XXX-XXX-XXX. Filled once for the whole sheet.", "AutoAccounting")
    from_letter = get_column_letter(_col(identity_headers, "from_date"))
    ws[f"{from_letter}2"].comment = Comment(
        "MM/DD/YYYY. One period applies to every form on this sheet.", "AutoAccounting")
    block_letter = get_column_letter(_col(entity_headers, "block"))
    ws[f"{block_letter}6"].comment = Comment(
        "'ewt' (Income Payments Subject to EWT) or 'mp' (Money Payments).", "AutoAccounting")
    name_letter = get_column_letter(_col(entity_headers, f"{var_prefix}_name"))
    ws[f"{name_letter}6"].comment = Comment(
        f"One row per {var_prefix}. Leave the name blank to add another line "
        f"item to the {var_prefix} above.", "AutoAccounting")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF rendering
# --------------------------------------------------------------------------- #
import fitz  # PyMuPDF

from app.forms import bir2307_layout as L


def _digits_only(s: str) -> str:
    return "".join(c for c in s if c.isdigit())


def _draw_centered(page: "fitz.Page", x: float, y: float, ch: str) -> None:
    if not ch:
        return
    w = fitz.get_text_length(ch, fontname=L.FONT_NAME, fontsize=L.FONT_SIZE)
    page.insert_text((x - w / 2, y), ch, fontname=L.FONT_NAME, fontsize=L.FONT_SIZE)


def _draw_per_cell(page: "fitz.Page", cells: list[float], y: float, value: str) -> None:
    for i, ch in enumerate(value):
        if i >= len(cells):
            break
        _draw_centered(page, cells[i], y, ch)


def _draw_text(page: "fitz.Page", spec: tuple[float, float, float], value: str) -> None:
    if not value:
        return
    x, y, _max_w = spec
    page.insert_text((x, y), value, fontname=L.FONT_NAME, fontsize=L.FONT_SIZE)


def _draw_col(page: "fitz.Page", col: tuple[float, float, str], y: float, value: str) -> None:
    if not value:
        return
    left, right, align = col
    if align == "right":
        w = fitz.get_text_length(value, fontname=L.FONT_NAME, fontsize=L.FONT_SIZE)
        x = right - 1.5 - w
    else:
        x = left + 1.5
    page.insert_text((x, y), value, fontname=L.FONT_NAME, fontsize=L.FONT_SIZE)


def _draw_date(page: "fitz.Page", cells: list[float], value: str) -> None:
    digits = _digits_only(value)
    if len(digits) != 8:
        return
    for i, ch in enumerate(digits):
        _draw_centered(page, cells[i], L.DATE_BASELINE_Y, ch)


def _draw_tin(page: "fitz.Page", cells: list[float], y: float, value: str) -> None:
    _draw_per_cell(page, cells, y, _digits_only(value))


def _draw_zip(page: "fitz.Page", y: float, value: str) -> None:
    _draw_per_cell(page, L.ZIP_CELLS, y, _digits_only(value))


def _fmt_amount(n: float) -> str:
    return f"{n:,.2f}"


def _fmt_amount_or_blank(n: float) -> str:
    return _fmt_amount(n) if n != 0 else ""


def _draw_description(page: "fitz.Page", y_top: float, text: str) -> None:
    if not text:
        return
    left, right, _ = L.COL_DESCRIPTION
    inner_w = (right - 1.5) - (left + 1.5)
    fontsize = L.FONT_SIZE
    while fontsize >= 5:
        if fitz.get_text_length(text, fontname=L.FONT_NAME, fontsize=fontsize) <= inner_w:
            break
        fontsize -= 0.5
    while fitz.get_text_length(text, fontname=L.FONT_NAME, fontsize=fontsize) > inner_w and len(text) > 1:
        text = text[:-1]
    page.insert_text((left + 1.5, y_top + L.LINE_TEXT_Y_OFFSET), text,
                     fontname=L.FONT_NAME, fontsize=fontsize)


def _draw_line_items(page: "fitz.Page", items: list[LineItem],
                     first_top_y: float, max_rows: int) -> None:
    for i, li in enumerate(items[:max_rows]):
        y_top = first_top_y + i * L.LINE_ROW_HEIGHT
        y = y_top + L.LINE_TEXT_Y_OFFSET
        _draw_description(page, y_top, li.description)
        _draw_col(page, L.COL_ATC,          y, li.atc)
        _draw_col(page, L.COL_M1,           y, _fmt_amount_or_blank(li.m1))
        _draw_col(page, L.COL_M2,           y, _fmt_amount_or_blank(li.m2))
        _draw_col(page, L.COL_M3,           y, _fmt_amount_or_blank(li.m3))
        _draw_col(page, L.COL_TOTAL,        y, _fmt_amount(li.total))
        _draw_col(page, L.COL_TAX_WITHHELD, y, _fmt_amount(li.tax_withheld))


def _fill_one_form(template_bytes: bytes, form: FilledForm) -> bytes:
    doc = fitz.open(stream=template_bytes, filetype="pdf")
    try:
        page = doc[0]
        _draw_date(page, L.DATE_CELLS_FROM, form.period_from)
        _draw_date(page, L.DATE_CELLS_TO,   form.period_to)
        _draw_tin(page, L.PAYEE_TIN_CELLS, L.PAYEE_TIN_BASELINE_Y, form.payee.tin)
        _draw_text(page, L.PAYEE_NAME, form.payee.name)
        _draw_text(page, L.PAYEE_REGISTERED_ADDRESS, form.payee.registered_address)
        _draw_zip(page, L.PAYEE_ZIP_BASELINE_Y, form.payee.zip_code)
        _draw_text(page, L.PAYEE_FOREIGN_ADDRESS, form.payee.foreign_address)
        _draw_tin(page, L.PAYOR_TIN_CELLS, L.PAYOR_TIN_BASELINE_Y, form.payor.tin)
        _draw_text(page, L.PAYOR_NAME, form.payor.name)
        _draw_text(page, L.PAYOR_REGISTERED_ADDRESS, form.payor.registered_address)
        _draw_zip(page, L.PAYOR_ZIP_BASELINE_Y, form.payor.zip_code)
        _draw_line_items(page, form.ewt, L.EWT_FIRST_ROW_TOP_Y, L.EWT_ROWS)
        _draw_line_items(page, form.money_payments, L.MP_FIRST_ROW_TOP_Y, L.MP_ROWS)
        return doc.tobytes()
    finally:
        doc.close()


def generate_merged_pdf(forms: list[FilledForm]) -> bytes:
    template_bytes = L.TEMPLATE_PATH.read_bytes()
    out = fitz.open()
    try:
        for form in forms:
            page_bytes = _fill_one_form(template_bytes, form)
            single = fitz.open(stream=page_bytes, filetype="pdf")
            try:
                out.insert_pdf(single)
            finally:
                single.close()
        return out.tobytes()
    finally:
        out.close()


# --------------------------------------------------------------------------- #
# E-signature stamping
# --------------------------------------------------------------------------- #
def apply_signature_to_pdf(
    pdf_bytes: bytes, signature_png: bytes, signer_name: str
) -> bytes:
    """Stamp a signature image + printed name onto the payor box of every page.

    The signature image is fitted (aspect-preserved) into the upper portion of
    ``L.PAYOR_SIGNATURE_BOX``; the printed ``signer_name`` is drawn centered just
    beneath it, inside the same signing strip. Applied to all pages so a
    multi-form document is signed once.
    """
    x0, y0, x1, y1 = L.PAYOR_SIGNATURE_BOX
    # Reserve the bottom ~9pt of the strip for the printed name; the image fills
    # the rest above it.
    name_band = 9.0
    img_rect = fitz.Rect(x0, y0, x1, max(y0 + 1, y1 - name_band))
    name_baseline = y1 - 1.5
    name = (signer_name or "").strip()

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        for page in doc:
            page.insert_image(img_rect, stream=signature_png, keep_proportion=True)
            if name:
                w = fitz.get_text_length(
                    name, fontname=L.FONT_NAME, fontsize=L.SIGNATURE_NAME_FONT_SIZE
                )
                cx = (x0 + x1) / 2
                page.insert_text(
                    (cx - w / 2, name_baseline), name,
                    fontname=L.FONT_NAME, fontsize=L.SIGNATURE_NAME_FONT_SIZE,
                )
        return doc.tobytes()
    finally:
        doc.close()


def bir2307_signature_placement(pdf_bytes: bytes) -> list[dict]:
    """Per-page signature marker boxes for the recipient's signing preview.

    BIR 2307 forms are signed at the fixed ``L.PAYOR_SIGNATURE_BOX`` on every
    page (see ``apply_signature_to_pdf``). We surface that same box as page
    fractions (top-left origin) so the signer sees exactly where their
    signature will land before they sign. Returned shape matches
    ``schemas.signature.SignaturePlacement`` (page, x, y, w, h).
    """
    x0, y0, x1, y1 = L.PAYOR_SIGNATURE_BOX
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        boxes: list[dict] = []
        for i, page in enumerate(doc, start=1):
            pw, ph = page.rect.width, page.rect.height
            if pw <= 0 or ph <= 0:
                continue
            boxes.append({
                "page": i,
                "x": x0 / pw,
                "y": y0 / ph,
                "w": (x1 - x0) / pw,
                "h": (y1 - y0) / ph,
            })
        return boxes
    finally:
        doc.close()


# Geometry for the custom-PDF signature strip (points). Sized relative to the
# last page, anchored to its bottom-left margin.
_CUSTOM_MARGIN = 40.0
_CUSTOM_BOX_WIDTH = 240.0
_CUSTOM_BOX_HEIGHT = 78.0
_CUSTOM_NAME_FONT_SIZE = 9.0
_CUSTOM_META_FONT_SIZE = 7.0


def _stamp_default_strip(page, signature_png: bytes, name: str) -> None:
    """Draw the fallback signature strip on the bottom-left of a page."""
    from datetime import datetime, timezone

    pw, ph = page.rect.width, page.rect.height
    box_w = min(_CUSTOM_BOX_WIDTH, pw - 2 * _CUSTOM_MARGIN)
    box_h = min(_CUSTOM_BOX_HEIGHT, ph - 2 * _CUSTOM_MARGIN)
    x0 = _CUSTOM_MARGIN
    y1 = ph - _CUSTOM_MARGIN
    x1 = x0 + box_w
    y0 = y1 - box_h

    # Bottom band holds the printed name + caption; the line sits above it; the
    # signature image fills the space above the line.
    caption_baseline = y1 - 2.0
    name_baseline = caption_baseline - (_CUSTOM_META_FONT_SIZE + 2.0)
    line_y = name_baseline - (_CUSTOM_NAME_FONT_SIZE + 2.0)
    img_rect = fitz.Rect(x0, y0, x1, max(y0 + 1, line_y - 3.0))

    page.insert_image(img_rect, stream=signature_png, keep_proportion=True)
    page.draw_line((x0, line_y), (x1, line_y), width=0.6)
    if name:
        page.insert_text(
            (x0, name_baseline), name,
            fontname=L.FONT_NAME, fontsize=_CUSTOM_NAME_FONT_SIZE,
        )
    signed_on = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    page.insert_text(
        (x0, caption_baseline), f"Electronically signed on {signed_on}",
        fontname=L.FONT_NAME, fontsize=_CUSTOM_META_FONT_SIZE,
    )


def _stamp_at_placement(page, signature_png: bytes, name: str, placement: dict) -> None:
    """Stamp the signature into a creator-defined box (fractions of the page)."""
    pw, ph = page.rect.width, page.rect.height
    # Clamp the normalized box so it stays on the page.
    fx = min(max(float(placement["x"]), 0.0), 1.0)
    fy = min(max(float(placement["y"]), 0.0), 1.0)
    fw = min(max(float(placement["w"]), 0.0), 1.0 - fx)
    fh = min(max(float(placement["h"]), 0.0), 1.0 - fy)
    x0, y0 = fx * pw, fy * ph
    x1, y1 = (fx + fw) * pw, (fy + fh) * ph

    # Reserve a small band at the bottom of the box for the printed name.
    name_band = _CUSTOM_NAME_FONT_SIZE + 3.0 if name else 0.0
    img_rect = fitz.Rect(x0, y0, x1, max(y0 + 1, y1 - name_band))
    page.insert_image(img_rect, stream=signature_png, keep_proportion=True)
    if name:
        w = fitz.get_text_length(name, fontname=L.FONT_NAME, fontsize=_CUSTOM_NAME_FONT_SIZE)
        cx = (x0 + x1) / 2
        page.insert_text(
            (cx - w / 2, y1 - 2.0), name,
            fontname=L.FONT_NAME, fontsize=_CUSTOM_NAME_FONT_SIZE,
        )


def apply_signature_to_custom_pdf(
    pdf_bytes: bytes,
    signature_png: bytes,
    signer_name: str,
    placement: list | dict | None = None,
) -> bytes:
    """Stamp a signature onto an uploaded PDF.

    ``placement`` may be:

    * a **list** of ``{page, x, y, w, h}`` (fractions, top-left origin, ``page``
      1-based); each listed page is stamped at its own box; unlisted pages are
      left untouched (per-page placement with include/exclude).
    * a **dict** ``{page, x, y, w, h}`` with ``page`` 1-based, or ``page=None``
      for every page (legacy single/all-page form).
    * ``None``: falls back to a self-contained strip on the last page only.

    The signature image is fitted into each box with the printed name beneath it.
    """
    name = (signer_name or "").strip()
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        if placement is None:
            _stamp_default_strip(doc[-1], signature_png, name)
        elif isinstance(placement, list):
            for item in placement:
                idx = max(1, min(int(item["page"]), doc.page_count)) - 1
                _stamp_at_placement(doc[idx], signature_png, name, item)
        elif placement.get("page") is None:
            for page in doc:
                _stamp_at_placement(page, signature_png, name, placement)
        else:
            idx = max(1, min(int(placement["page"]), doc.page_count)) - 1
            _stamp_at_placement(doc[idx], signature_png, name, placement)
        return doc.tobytes()
    finally:
        doc.close()
