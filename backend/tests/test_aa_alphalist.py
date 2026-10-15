"""1604-C alphalist, 2316 values, year-end adjustment, payroll sheet import (pure)."""

from datetime import date
from decimal import Decimal

import fitz

from app.services.accounting.compute.alphalist import (
    MonthRemittance, build_alphalist, compute_employee, employee_from_json,
)
from app.services.accounting.payroll_import import apply_mapping, parse_sheet, suggest_mapping
from app.services.accounting.pdf_2316 import build_2316_zip, fill_2316

D = Decimal


def emp(**data):
    base = {"tin": "123-456-789", "last_name": "Dela Cruz", "first_name": "Juan", "status": "R"}
    base.update(data)
    return employee_from_json(data.pop("id", "e1") if "id" in data else "e1", base)


def test_year_end_adjustment_under_withheld_goes_to_december():
    # taxable 600,000 (2026): 22,500 + 20% × 200,000 = 62,500 due; 55,000 withheld Jan–Nov → 7,500 in December
    c = compute_employee(emp(t_basic="560000", t_13th="40000", nt_13th="90000", nt_deminimis="10000",
                             nt_contributions="24000", tax_withheld_jan_nov="55000"), 2026)
    a = c.alpha
    assert c.schedule == 1
    assert (a["7f"], a["7j"], a["7a"]) == (D("124000.00"), D("600000.00"), D("724000.00"))
    assert (a["7g"], a["7h"], a["7i"]) == (D("560000.00"), D("40000.00"), D("0.00"))
    assert (a["14"], a["17a"], a["17b"], a["18"], a["19"]) == (D("62500.00"), D("7500.00"), D("0.00"), D("62500.00"), "Y")
    assert (c.form["19"], c.form["21"], c.form["24"], c.form["25A"], c.form["28"]) == (
        D("724000.00"), D("600000.00"), D("62500.00"), D("62500.00"), D("62500.00"))


def test_over_withholding_is_refunded_and_previous_employer_blocks_substituted_filing():
    c = compute_employee(emp(t_basic="300000", tax_withheld_jan_nov="9000", prev_tin="111-222-333",
                             prev_t_basic="100000", tax_withheld_prev="4000"), 2026)
    # 400,000 taxable: 15% × (400,000 − 250,000) = 22,500 due; 13,000 withheld → 9,500 in December
    assert c.alpha["13"] == D("400000.00") and c.alpha["14"] == D("22500.00")
    assert c.alpha["17a"] == D("9500.00") and c.alpha["19"] == "N"
    over = compute_employee(emp(t_basic="260000", tax_withheld_jan_nov="5000"), 2026)
    assert (over.alpha["14"], over.alpha["17b"], over.form["25A"]) == (D("1500.00"), D("3500.00"), D("1500.00"))


def test_minimum_wage_earner_goes_to_schedule_2_and_13th_month_over_ceiling_warns():
    c = compute_employee(emp(mwe=True, region="NCR", smw_day="645", smw_month="16770", nt_basic="201240",
                             nt_overtime="5000", nt_13th="95000"), 2026)
    assert c.schedule == 2 and c.alpha["7f"] == D("201240.00") and c.alpha["7h"] == D("5000.00")
    assert c.alpha["14"] == D("0.00")
    assert any("above ₱90,000" in w for w in c.warnings)


def test_alphalist_totals_remittances_and_checks():
    employees = [emp(id="a", t_basic="300000", tax_withheld_jan_nov="7500"),
                 emp(id="b", last_name="Aquino", first_name="Ana", tin="222-333-444", t_basic="250000")]
    rem = [MonthRemittance(m, "2026-02-10", D("45833.33"), D("625"), D("0"), D("0"), D("625"), "filed") for m in range(1, 13)]
    tabs, computed, warnings = build_alphalist(employees, 2026, rem)
    rows = tabs["Schedule1"]
    assert [r["last"] for r in rows] == ["AQUINO", "DELA CRUZ", "TOTALS"] and rows[0]["seq"] == "1"
    assert rows[-1]["7j"] == D("300000.00") and rows[-1]["7e"] == D("250000.00")  # ₱250k or less is exempt
    assert tabs["Remittances"][-1]["withheld"] == D("7500.00")
    withheld = next(r for r in tabs["Checks"] if r["key"] == "withheld")
    assert withheld["difference"] == D("0.00")
    gross = next(r for r in tabs["Checks"] if r["key"] == "gross")
    assert gross["difference"] == D("550000.00") - D("549999.96")
    assert any("differ by" in w for w in warnings)


def test_import_suggests_mapping_and_reads_rows(tmp_path):
    csv = ("Payroll summary 2026\n"
           "Employee Name,TIN,Date Hired,Basic Salary,13th Month,SSS/PhilHealth/Pag-IBIG,De Minimis,Tax Withheld,MWE\n"
           "\"Dela Cruz, Juan P.\",123-456-789,01/15/2026,\"300,000.00\",25000,\"18,000\",5000,7500,\n"
           "\"Santos, Maria\",222-333-444,2026-03-01,201240,16770,12000,0,0,Y\n"
           "TOTAL,,,501240,41770,30000,5000,7500,\n").encode()
    headers, rows = parse_sheet(csv, "payroll.csv")
    assert headers[0] == "Employee Name" and len(rows) == 3
    mapping = suggest_mapping(headers)
    assert mapping == {"Employee Name": "full_name", "TIN": "tin", "Date Hired": "employed_from",
                       "Basic Salary": "t_basic", "13th Month": "nt_13th", "SSS/PhilHealth/Pag-IBIG": "nt_contributions",
                       "De Minimis": "nt_deminimis", "Tax Withheld": "tax_withheld_jan_nov", "MWE": "mwe"}
    records, errors = apply_mapping(headers, rows, mapping)
    assert errors == [] and len(records) == 2
    assert records[0]["last_name"] == "Dela Cruz" and records[0]["first_name"] == "Juan"
    assert records[0]["middle_name"] == "P." and records[0]["t_basic"] == "300000.00"
    assert records[0]["employed_from"] == "2026-01-15" and records[1]["mwe"] is True
    assert suggest_mapping(headers, {"Basic Salary": "nt_basic"})["Basic Salary"] == "nt_basic"


def test_2316_pdf_fills_the_official_form():
    c = compute_employee(emp(t_basic="560000", nt_13th="90000", tax_withheld_jan_nov="55000",
                             employed_from="2026-01-15"), 2026)
    doc = fitz.open()
    fill_2316(doc, c, {"tin": "010234567000", "name": "Bayanihan Hardware Trading Corp.", "address": "Manila"})
    text = doc[0].get_text()
    assert "DELA CRUZ, JUAN" in text and "560,000.00" in text and "650,000.00" in text
    assert "BAYANIHAN HARDWARE" in text
    data = build_2316_zip([c], {"tin": "010234567000", "name": "X"}, 2026)
    assert data[:2] == b"PK"


def test_compensation_of_250k_or_less_is_reported_as_non_taxable():
    c = compute_employee(emp(t_basic="240000", t_13th="5000", tax_withheld_jan_nov="1200"), 2026)
    assert (c.alpha["7e"], c.alpha["7b"], c.alpha["7j"], c.alpha["14"]) == (
        D("240000.00"), D("5000.00"), D("0.00"), D("0.00"))
    assert (c.form["29"], c.form["52"], c.alpha["17b"]) == (D("240000.00"), D("0.00"), D("1200.00"))


def test_imported_text_never_becomes_a_formula_in_the_workbook():
    import io as _io
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from app.routers.accounting_employees import alphalist_workbook
    tabs, _, warnings = build_alphalist([emp(last_name="=HYPERLINK(\"http://x\")", first_name="+1")], 2026, [])
    company = SimpleNamespace(registered_name="X Corp", tin="123456789", branch_code="00000")
    wb = load_workbook(_io.BytesIO(alphalist_workbook(company, 2026, tabs, warnings)))
    values = [c.value for row in wb["Schedule 1"].iter_rows() for c in row if isinstance(c.value, str)]
    assert not any(v.startswith(("=", "+")) for v in values)


def test_oversized_workbooks_and_wide_sheets_are_bounded():
    import io as _io
    import zipfile as _zip
    import pytest
    buf = _io.BytesIO()
    with _zip.ZipFile(buf, "w", _zip.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (61 * 1024 * 1024))
    with pytest.raises(ValueError, match="too large"):
        parse_sheet(buf.getvalue(), "bomb.xlsx")
    wide = ("A,B,C," + ",".join(f"c{i}" for i in range(500)) + "\n").encode()
    headers, _ = parse_sheet(wide, "wide.csv")
    assert len(headers) == 120
