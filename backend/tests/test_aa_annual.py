from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.services.accounting.compute.annual import (
    CarryIn, annual_form, build_annual, carry_in_from_snapshot, ordinary_schedule,
)
from app.services.accounting.compute.annual_inputs import AnnualInputs, QuarterSource
from app.services.accounting.compute.depreciation import Asset
from app.services.accounting.compute.inputs import PeriodData, PeriodInputs
from app.services.accounting.diff import snapshot
from app.services.accounting.render_xlsx import write_workbook
from tests.aa_factories import company, expense, line, purchase, sale

D = Decimal


def P(year, quarter, *lines, **inputs):
    return PeriodData(year, quarter, tuple(lines), PeriodInputs(**inputs))


def _machine(cost="60000", in_service=date(2026, 1, 15), life=60):
    return Asset("a1", "d1", "Sewing machine", "SI-9", D(cost), in_service, life)


def _year(year, *extra_q1, sales="250000", vat="30000"):
    gross = str(D(sales) + D(vat))
    return [P(year, 1, sale(vatable=sales, vat=vat, gross=gross, doc_date=date(year, 2, 1)), *extra_q1,
              opening_inventory=D("0")),
            *[P(year, q, sale(vatable=sales, vat=vat, gross=gross, doc_date=date(year, q * 3 - 1, 1)))
              for q in (2, 3, 4)]]


def _sources(*payables):
    return {q: QuarterSource(q, D(p), 1, True, 1) for q, p in enumerate(payables, start=1)}


def test_annual_form_by_taxpayer():
    assert annual_form(company()) == "1702RT"
    assert annual_form(company(taxpayer_type="individual")) == "1701"
    assert annual_form(company(taxpayer_type="individual", deduction_method="osd")) == "1701A"
    assert annual_form(company(taxpayer_type="individual", vat_registered=False,
                               income_tax_option="eight_percent")) == "1701A"


def test_corporation_itemized_profit():
    periods = _year(2026, purchase(gross="448000", vat="48000", input_vat_valid=True),
                    expense(account="rent", gross="134400", vat="14400", input_vat_valid=True))
    r = build_annual(company(), periods, 2026, AnnualInputs(ending_inventory=D("100000")), assets=[_machine()],
                     quarter_sources=_sources("10000", "20000", "30000"))
    f = r.forms["1702RT"]
    assert (f["27"], f["30"], f["31"], f["33"]) == (D("1000000"), D("300000"), D("700000"), D("700000"))
    assert (f["S1_5"], f["S1_11"], f["S1_18"], f["34"], f["37"]) == (
        D("12000"), D("120000"), D("132000"), D("132000"), D("132000"))
    assert (f["39"], f["40"], f["41"], f["42"], f["43"]) == (D("568000"), D("25"), D("142000"), D("0"), D("142000"))
    assert (f["46"], f["55"], f["56"], f["14"], f["21"]) == (
        D("60000"), D("60000"), D("82000"), D("142000"), D("82000"))
    assert r.ledgers["Depreciation"][0]["this_year"] == D("12000.00")
    assert [row["source"] for row in r.ledgers["QuarterCredits"]] == ["filed v1"] * 3 + ["this return"]
    assert r.forms["COGS"]["cost_of_sales"] == D("300000.00")
    assert r.forms["Expenses"]["depreciation"] == D("12000.00")
    assert any("MCIT" in w and "Q3" in w for w in r.warnings)
    assert r.forms["1702RT"]["OVERPAYMENT_OPTION"] == ""


def test_loss_year_then_profit_uses_nolco():
    services = company(business_type="services")
    loss = build_annual(services, [P(2025, 1, sale(vatable="100000"), expense(account="rent", gross="300000"))],
                        2025, AnnualInputs())
    assert loss.forms["1702RT"]["39"] == D("-200000") and loss.forms["1702RT"]["S3_3"] == D("200000")
    assert loss.carry_out["nolco"] == [{"year": 2025, "amount": "200000.00", "used": "0.00"}]
    assert loss.forms["1702RT"]["S3_7_YEAR"] == "2025" and loss.forms["1702RT"]["S3_7_A"] == D("200000")

    carry = carry_in_from_snapshot(loss.carry_out)
    profit = build_annual(services, [P(2026, 1, sale(vatable="500000"), expense(account="rent", gross="100000"))],
                          2026, AnnualInputs(), carry_in=carry)
    f = profit.forms["1702RT"]
    assert (f["36"], f["37"], f["39"], f["41"], f["S3_8"]) == (
        D("200000"), D("300000"), D("200000"), D("50000"), D("200000"))
    assert f["S3_4_YEAR"] == "2025" and f["S3_4_D"] == D("200000") and f["S3_4_E"] == D("0")
    assert profit.carry_out["nolco"] == []
    assert [row["year"] for row in profit.ledgers["NOLCO"]] == ["2025"]


def test_mcit_year_then_rcit_year():
    corp = company(business_type="services", operations_start_year=2020)
    y1 = build_annual(corp, [P(2025, 1, sale(vatable="1000000"), expense(account="rent", gross="990000"))],
                      2025, AnnualInputs())
    f1 = y1.forms["1702RT"]
    assert (f1["41"], f1["42"], f1["43"]) == (D("2500"), D("20000"), D("20000"))
    assert y1.carry_out["excess_mcit"] == [{"year": 2025, "amount": "17500.00", "used": "0.00",
                                            "normal_tax": "2500.00", "mcit": "20000.00"}]
    y2 = build_annual(corp, [P(2026, 1, sale(vatable="1000000"), expense(account="rent", gross="500000"))],
                      2026, AnnualInputs(), carry_in=carry_in_from_snapshot(y1.carry_out))
    f2 = y2.forms["1702RT"]
    assert (f2["41"], f2["42"], f2["47"], f2["S4_4"], f2["56"]) == (
        D("125000"), D("20000"), D("17500"), D("17500"), D("107500"))
    assert f2["S4_1_YEAR"] == "2025" and f2["S4_1_A"] == D("2500") and f2["S4_1_B"] == D("20000")
    # a smaller RCIT margin caps the credit at RCIT − MCIT
    y3 = build_annual(corp, [P(2026, 1, sale(vatable="1000000"), expense(account="rent", gross="900000"))],
                      2026, AnnualInputs(), carry_in=carry_in_from_snapshot(y1.carry_out))
    assert (y3.forms["1702RT"]["41"], y3.forms["1702RT"]["47"]) == (D("25000"), D("5000"))


def test_mcit_before_the_fourth_year_and_without_start_year():
    young = company(business_type="services", operations_start_year=2023)
    r = build_annual(young, [P(2026, 1, sale(vatable="1000000"), expense(account="rent", gross="990000"))],
                     2026, AnnualInputs())
    assert r.forms["1702RT"]["42"] == D("0")
    q3 = build_annual(company(business_type="services"),
                      [P(2026, 3, sale(vatable="1000000"), expense(account="rent", gross="990000"), mcit_applies=True)],
                      2026, AnnualInputs())
    assert q3.forms["1702RT"]["42"] == D("20000") and any("Q3" in w for w in q3.warnings)


def test_cit_20_percent_only_up_to_5m():
    services = company(business_type="services")
    small = build_annual(services, [P(2026, 1, sale(vatable="1000000"))], 2026, AnnualInputs(cit_20_percent=True))
    assert (small.forms["1702RT"]["40"], small.forms["1702RT"]["41"]) == (D("20"), D("200000"))
    big = build_annual(services, [P(2026, 1, sale(vatable="6000000"))], 2026, AnnualInputs(cit_20_percent=True))
    assert big.forms["1702RT"]["40"] == D("25") and any("₱5,000,000" in w for w in big.warnings)


def test_individual_osd_and_eight_percent_1701a():
    osd = company(taxpayer_type="individual", vat_registered=False, deduction_method="osd", business_type="services")
    r = build_annual(osd, [P(2026, 1, sale(gross="1000000"))], 2026, AnnualInputs())
    f = r.forms["1701A"]
    assert (f["36"], f["39"], f["40"], f["45"], f["46"], f["20"], f["29"]) == (
        D("1000000"), D("400000"), D("600000"), D("600000"), D("62500"), D("62500"), D("62500"))
    assert "56" not in f and not any("compensation income" in w.lower() for w in r.warnings)

    eight = company(taxpayer_type="individual", vat_registered=False, income_tax_option="eight_percent",
                    business_type="services")
    e = build_annual(eight, [P(2026, 1, sale(gross="3500000"))], 2026, AnnualInputs())
    g = e.forms["1701A"]
    assert (g["53"], g["54"], g["55"], g["56"], g["20"]) == (
        D("3500000"), D("250000"), D("3250000"), D("260000"), D("260000"))
    assert "46" not in g and any("₱3,000,000" in w for w in e.warnings)


def test_individual_itemized_1701():
    person = company(taxpayer_type="individual", vat_registered=False, business_type="services")
    r = build_annual(person, [P(2026, 1, sale(gross="1000000"), expense(account="rent", gross="200000"))], 2026,
                     AnnualInputs(), quarter_sources=_sources("50000"))
    f = r.forms["1701"]
    assert (f["S3_8"], f["S3_12"], f["S3_13"], f["S3_16"], f["S3_18"], f["S3_23"], f["S3_25"]) == (
        D("1000000"), D("1000000"), D("200000"), D("200000"), D("800000"), D("800000"), D("102500"))
    assert (f["S4_11"], f["S4_18"], f["P6_5"], f["P7_2"], f["22"], f["31"]) == (
        D("200000"), D("200000"), D("102500"), D("50000"), D("102500"), D("52500"))


def test_june_fiscal_year():
    june = company(business_type="services", fiscal_year_end_month=6)
    asset = _machine(cost="12000", in_service=date(2025, 12, 5), life=12)
    r = build_annual(june, [P(2026, 1, sale(vatable="100000", doc_date=date(2025, 8, 1))),
                            P(2026, 4, sale(vatable="100000", doc_date=date(2026, 5, 1)))],
                     2026, AnnualInputs(), assets=[asset])
    assert r.forms["1702RT"]["S1_5"] == D("7000") and r.forms["1702RT"]["27"] == D("200000")


def test_reconciliation_and_balance_sheet():
    periods = _year(2026, purchase(gross="448000", vat="48000", input_vat_valid=True),
                    expense(account="rent", gross="134400", vat="14400", input_vat_valid=True))
    inputs = AnnualInputs(ending_inventory=D("100000"), book_net_income=D("550000"),
                          nondeductible_expenses=D("20000"), other_reconciling=D("-2000"),
                          current_assets=D("1000000"), current_liabilities=D("400000"), equity=D("500000"))
    r = build_annual(company(), periods, 2026, inputs, assets=[_machine()])
    f = r.forms["1702RT"]
    assert (f["S5_1"], f["S5_2"], f["S5_4"], f["S5_7"], f["S5_9"], f["S5_10"]) == (
        D("550000"), D("20000"), D("570000"), D("2000"), D("2000"), D("568000"))
    assert not any("Schedule V" in w for w in r.warnings)
    bs = r.forms["BalanceSheet"]
    assert bs["ppe_net"] == D("48000.00") and bs["total_assets"] == D("1048000.00")
    assert any("balance" in w.lower() for w in r.warnings)
    off = build_annual(company(), periods, 2026, AnnualInputs(ending_inventory=D("100000"),
                                                              book_net_income=D("500000")), assets=[_machine()])
    assert any("Schedule V" in w for w in off.warnings)
    blank = build_annual(company(), periods, 2026, AnnualInputs(ending_inventory=D("100000")))
    assert "S5_1" not in blank.forms["1702RT"] and "BalanceSheet" not in blank.forms


def test_overpayment_carried_unless_refund():
    services = company(business_type="services")
    periods = [P(2026, 1, sale(vatable="100000"))]
    r = build_annual(services, periods, 2026, AnnualInputs(), quarter_sources=_sources("50000"))
    assert r.forms["1702RT"]["56"] == D("-25000") and r.forms["1702RT"]["OVERPAYMENT_OPTION"] == "carry_over"
    assert r.carry_out["excess_credits"] == "25000"
    assert carry_in_from_snapshot(r.carry_out).excess_credits == D("25000")
    refund = build_annual(services, periods, 2026, AnnualInputs(refund_overpayment=True),
                          quarter_sources=_sources("50000"))
    assert refund.forms["1702RT"]["OVERPAYMENT_OPTION"] == "refund" and refund.carry_out["excess_credits"] is None
    nxt = build_annual(services, [P(2027, 1, sale(vatable="100000"))], 2027, AnnualInputs(),
                       carry_in=carry_in_from_snapshot(r.carry_out))
    assert nxt.forms["1702RT"]["44"] == D("25000")


def test_ordinary_schedule_maps_accounts():
    from app.services.accounting.compute.deductions import quarter_expenses

    lines = [expense(account="rent", gross="1000"), expense(account="utilities", gross="200"),
             expense(account="professional_fees", gross="300"), line(category="payroll", gross="500")]
    b, _, _ = quarter_expenses(lines, company(business_type="services"), D("0"), D("0"), D("100000"), D("70"))
    s = ordinary_schedule(b, D("50"))
    assert s == {"5": D("50"), "6": D("0"), "11": D("1000"), "13": D("500"), "14": D("70"), "15": D("0"),
                 "16": D("0"), "17B": D("300"), "17D": D("200"), "18": D("2120")}


def test_snapshot_through_write_workbook(tmp_path):
    periods = _year(2026, purchase(gross="448000", vat="48000", input_vat_valid=True))
    r = build_annual(company(), periods, 2026, AnnualInputs(ending_inventory=D("100000")), assets=[_machine()],
                     carry_in=CarryIn())
    snap = snapshot(r, company())
    path = tmp_path / "annual.xlsx"
    write_workbook(path, header={"company": "Juan", "tin": "123-456-789-00000", "period": "FY 2026 · Annual",
                                 "generated": "2026-10-06 10:00"},
                   snapshot=snap, warnings=r.warnings, review_open=0, due={"1702RT": "2027-04-15"}, version=1)
    wb = load_workbook(path)
    assert {"Summary", "1702RT", "COGS", "Expenses", "Depreciation", "QuarterCredits", "Transactions"} <= set(
        wb.sheetnames)
    ws = wb["1702RT"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "21")
    assert ws.cell(row, 3).number_format == "#,##0"
    summary = [c.value for c in wb["Summary"]["A"]]
    assert any(v and "1702-RT" in str(v) for v in summary)


def test_year_cells_stay_text_in_the_workbook(tmp_path):
    path = tmp_path / "y.xlsx"
    write_workbook(path, header={"company": "J", "tin": "1", "period": "FY 2026 · Annual", "generated": "now"},
                   snapshot={"company": {}, "forms": {"1702RT": {"S3_4_YEAR": "2024", "21": "5"}}, "ledgers": {}},
                   warnings=[], review_open=0, due={}, version=1)
    ws = load_workbook(path)["1702RT"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "S3_4_YEAR")
    assert ws.cell(row, 3).value == "2024"


def test_books_starting_mid_year_take_the_earlier_months_from_the_annual_inputs():
    gross = "280000"
    periods = [P(2026, 3, sale(vatable="250000", vat="30000", gross=gross, doc_date=date(2026, 8, 1)),
                 opening_inventory=D("0")),
               P(2026, 4, sale(vatable="250000", vat="30000", gross=gross, doc_date=date(2026, 11, 1)))]
    bare = build_annual(company(), periods, 2026, AnnualInputs(ending_inventory=D("0")), books_start=date(2026, 7, 1))
    assert any("aren't on AutoAccounting" in w for w in bare.warnings)
    assert not any(w.startswith("No documents for") for w in bare.warnings)

    inputs = AnnualInputs(ending_inventory=D("0"), pre_books_sales=D("400000"), pre_books_cost_of_sales=D("100000"),
                          pre_books_deductions=D("50000"), pre_books_cwt=D("4000"))
    r = build_annual(company(), periods, 2026, inputs, books_start=date(2026, 7, 1))
    f = r.forms["1702RT"]
    assert (f["27"], f["30"], f["S1_17D"], f["48"]) == (D("900000"), D("100000"), D("50000"), D("4000"))
    assert not any("aren't on AutoAccounting" in w for w in r.warnings)


def _mixed(**kw):
    return company(taxpayer_type="individual", vat_registered=False, business_type="services",
                   has_compensation_income=True, **kw)


_SALARY = dict(compensation_income=D("600000"), compensation_nontaxable=D("90000"),
               compensation_tax_withheld=D("40000"))


def test_mixed_income_files_1701_with_salary_and_business_at_graduated_rates():
    r = build_annual(_mixed(), [P(2026, 1, sale(gross="1000000"), expense(gross="400000"))], 2026,
                     AnnualInputs(**_SALARY))
    assert r.form == "1701" and annual_form(_mixed(deduction_method="osd")) == "1701"
    f = r.forms["1701"]
    assert (f["S2_4"], f["S2_5"], f["S2_6"]) == (D("600000"), D("90000"), D("510000"))
    # business 1,000,000 − 400,000 = 600,000; with salary 1,110,000 → 102,500 + 25% × 310,000
    assert (f["S3_23"], f["S3_24"], f["S3_25"], f["P7_5"], f["31"]) == (
        D("600000"), D("1110000"), D("180000"), D("40000"), D("140000"))


def test_mixed_income_osd_takes_40_percent_of_sales():
    f = build_annual(_mixed(deduction_method="osd"), [P(2026, 1, sale(gross="1000000"), expense(gross="400000"))],
                     2026, AnnualInputs(**_SALARY)).forms["1701"]
    assert (f["S3_17"], f["S3_18"], f["S3_24"], f["S3_25"]) == (D("400000"), D("600000"), D("1110000"), D("180000"))
    assert "S4_18" not in f and "S6_1" not in f


def test_mixed_income_eight_percent_uses_schedule_3b_without_the_250k():
    f = build_annual(_mixed(income_tax_option="eight_percent"), [P(2026, 1, sale(gross="1000000"))], 2026,
                     AnnualInputs(**_SALARY)).forms["1701"]
    # salary 510,000 → 22,500 + 20% × 110,000 = 44,500; business 8% × 1,000,000 = 80,000
    assert (f["S2_7"], f["S3_29"], f["S3_31"], f["S3_32"], f["22"], f["31"]) == (
        D("44500"), D("0"), D("80000"), D("124500"), D("124500"), D("84500"))
    assert "S3_25" not in f


def test_mixed_income_without_a_2316_warns():
    r = build_annual(_mixed(), [P(2026, 1, sale(gross="1000000"))], 2026, AnnualInputs())
    assert any("BIR 2316" in w for w in r.warnings)
