from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.services.accounting.compute.due_dates import due_dates
from app.services.accounting.compute.engine import build_quarter, expected_forms
from app.services.accounting.compute.inputs import (
    PayrollMonth,
    PeriodData,
    PeriodInputs,
)
from app.services.accounting.diff import snapshot
from app.services.accounting.render_xlsx import write_workbook
from tests.aa_factories import company, expense, line, sale

D = Decimal


def P(year, quarter, *lines, **inputs):
    inputs.setdefault("opening_inventory", D("0"))
    inputs.setdefault("ending_inventory", D("0"))
    return PeriodData(year, quarter, tuple(lines), PeriodInputs(**inputs))


def _rent():
    return expense(account="rent", counterparty_name="Ayala Land Inc.", counterparty_tin="111222333", gross="11200",
                   vat="1200", vatable="10000", withholding="500", input_vat_valid=True, doc_date=date(2026, 7, 5))


def test_withholding_agent_gets_0619e_1601eq_and_qap():
    c = company(is_withholding_agent=True)
    rent = _rent()
    r = build_quarter(c, [P(2026, 3, rent)])
    assert {"0619E-M1", "0619E-M2", "1601EQ"} <= set(r.forms)
    assert r.forms["0619E-M1"]["18"] == D("500.00") and r.forms["0619E-M2"]["18"] == D("0.00")
    assert r.forms["1601EQ"]["S1_ATC"] == "WC100" and r.forms["1601EQ"]["26"] == D("0.00")
    assert r.ledgers["QAP"][0]["tax"] == D("500.00")
    tx = r.ledgers["Transactions"][0]
    assert tx["ewt_atc"] == "WC100" and tx["ewt_due"] == D("500.00")
    assert expected_forms(c, 3) == ["2550Q", "1702Q", "0619E-M1", "0619E-M2", "1601EQ"]


def test_agent_with_no_payments_still_files_zero_returns():
    r = build_quarter(company(is_withholding_agent=True), [P(2026, 3, sale(vatable="1000", vat="120"))])
    assert r.forms["0619E-M1"]["18"] == D("0.00") and r.forms["1601EQ"]["31"] == D("0.00")
    assert "QAP" not in r.ledgers


def test_non_agent_has_no_withholding_forms_but_transactions_show_blank_atc():
    r = build_quarter(company(), [P(2026, 3, _rent())])
    assert not any(form.startswith(("0619E", "1601")) for form in r.forms)
    assert r.ledgers["Transactions"][0]["ewt_atc"] == "" and r.ledgers["Transactions"][0]["ewt_due"] == D("0.00")


def test_fiscal_year_not_on_calendar_quarters_skips_ewt_returns():
    c = company(is_withholding_agent=True, fiscal_year_end_month=5)
    r = build_quarter(c, [P(2026, 1)])
    assert "1601EQ" not in r.forms and any("calendar quarters" in w for w in r.warnings)
    assert "1601EQ" not in expected_forms(c, 1)


def test_employer_gets_three_1601c_months():
    c = company(has_employees=True)
    payroll = line(category="payroll", account="salaries", gross="100000", withholding="8000",
                   doc_date=date(2026, 8, 31))
    months = (PayrollMonth(), PayrollMonth(contributions=D("5000")), PayrollMonth())
    r = build_quarter(c, [P(2026, 3, payroll, payroll=months)])
    assert r.forms["1601C-M2"]["14"] == D("100000.00") and r.forms["1601C-M2"]["22"] == D("95000.00")
    assert r.forms["1601C-M2"]["36"] == D("8000.00") and r.forms["1601C-M1"]["36"] == D("0.00")
    assert expected_forms(c, 4)[-3:] == ["1601C-M1", "1601C-M2", "1601C-M3"]


def test_payroll_without_employee_flag_still_files_1601c_with_warning():
    payroll = line(category="payroll", account="salaries", gross="1000", doc_date=date(2026, 7, 15))
    r = build_quarter(company(), [P(2026, 3, payroll)])
    assert "1601C-M1" in r.forms and any("having employees" in w for w in r.warnings)


def test_due_dates_for_withholding_forms():
    days = due_dates(["0619E-M1", "0619E-M2", "1601EQ", "1601C-M3"], 2026, 3, 12)
    assert days["0619E-M1"] == date(2026, 8, 10)
    assert days["0619E-M2"] == date(2026, 9, 10)
    assert days["1601EQ"] == date(2026, 11, 2)     # Oct 31 is a Saturday
    assert days["1601C-M3"] == date(2026, 10, 12)  # Oct 10 is a Saturday
    assert due_dates(["1601C-M3"], 2026, 4, 12)["1601C-M3"] == date(2027, 1, 15)


def test_workbook_has_withholding_tabs(tmp_path):
    c = company(is_withholding_agent=True, has_employees=True)
    r = build_quarter(c, [P(2026, 3, _rent())])
    path = tmp_path / "w.xlsx"
    write_workbook(path, header={"company": "Juan", "tin": "1", "period": "2026 Q3", "generated": "now"},
                   snapshot=snapshot(r, c), warnings=r.warnings, review_open=0, due={}, version=1)
    names = load_workbook(path).sheetnames
    assert {"0619E-M1", "0619E-M2", "1601EQ", "1601C-M1", "QAP"} <= set(names)


def test_payroll_documents_put_1601c_in_expected_forms():
    assert "1601C-M1" not in expected_forms(company(), 3)
    assert expected_forms(company(), 3, has_payroll=True)[-3:] == ["1601C-M1", "1601C-M2", "1601C-M3"]
    payroll = line(category="payroll", account="salaries", gross="1000", doc_date=date(2026, 7, 15))
    r = build_quarter(company(), [P(2026, 3, payroll)])
    assert any("Has employees" in w and "non-taxable" in w for w in r.warnings)


def test_more_than_30_atcs_warns_that_the_tab_is_cut():
    from app.services.accounting.compute import withholding

    c = company(is_withholding_agent=True)
    lines = [expense(account="rent", ewt_atc=code, gross="1000") for code in ("WI100", "WC100", "WI120", "WC120")]
    original = withholding.MAX_SCHEDULE_ROWS
    withholding.MAX_SCHEDULE_ROWS = 3
    try:
        r = build_quarter(c, [P(2026, 3, *lines)])
    finally:
        withholding.MAX_SCHEDULE_ROWS = original
    assert "S4_ATC" not in r.forms["1601EQ"] and len(r.ledgers["QAP"]) == 4
    assert any("first 3" in w and "QAP" in w for w in r.warnings)


def test_filed_0619e_feeds_the_1601eq_and_calculated_months_warn():
    c = company(is_withholding_agent=True)
    r = build_quarter(c, [P(2026, 3, _rent())], filed_remittances={1: (D("400.00"), 2)})
    assert r.forms["1601EQ"]["20"] == D("400.00") and r.forms["1601EQ"]["_SRC_20"] == "filed:v2"
    assert r.forms["1601EQ"]["26"] == D("100.00")   # 500 withheld − 400 remitted
    assert any("0619-E M2 isn't marked filed" in w for w in r.warnings)
    assert not any("0619-E M1 isn't marked filed" in w for w in r.warnings)
