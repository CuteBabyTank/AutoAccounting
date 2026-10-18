"""Import routes, pure parts: options parsing, templates (formula-safe, round-trip through the importers)."""

import io
from datetime import date

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.routers.accounting_books import entry_problem
from app.routers.accounting_imports import chart_refs, chosen_columns, parse_options, template_bytes
from app.schemas.accounting_books import JournalEntryIn
from app.services.accounting.books.chart import Chart, standard_chart
from app.services.accounting.import_rows import build_journals, build_open_items, build_tb
from app.services.accounting.payroll_import import parse_sheet

CHART = Chart(standard_chart("corporation"))
START = date(2026, 1, 1)


def test_import_options_parse():
    assert parse_options("").columns is None
    opts = parse_options('{"columns": {"debit": "Dr"}, "accounts": {"|bank": "1020"}, "default_kind": "ar"}')
    assert opts.columns == {"debit": "Dr"} and opts.accounts == {"|bank": "1020"} and opts.default_kind == "ar"
    for bad in ("{", "[]", '{"default_kind": "xx"}'):
        with pytest.raises(HTTPException) as exc:
            parse_options(bad)
        assert exc.value.status_code == 400


def test_template_xlsx_keeps_numbers_and_guards_text():
    wb = load_workbook(io.BytesIO(template_bytes("journal-entries", "xlsx")))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0][0] == "Journal no."
    assert rows[-1][4] == -1500.0          # a negative amount is a number, not "'-1500.00" text


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_templates_import_cleanly(fmt):
    refs = chart_refs(CHART)
    headers, rows = parse_sheet(template_bytes("opening-tb", fmt), f"t.{fmt}")
    tb = build_tb(headers, rows, chosen_columns("opening-tb", headers, parse_options("")), refs, {},
                  books_start=START)
    assert tb.ready and not tb.warnings, (tb.errors, tb.warnings)
    headers, rows = parse_sheet(template_bytes("open-items", fmt), f"t.{fmt}")
    oi = build_open_items(headers, rows, chosen_columns("open-items", headers, parse_options("")), START)
    assert oi.ready and len(oi.items) == 2, oi.errors
    headers, rows = parse_sheet(template_bytes("journal-entries", fmt), f"t.{fmt}")
    je = build_journals(headers, rows, chosen_columns("journal-entries", headers, parse_options("")), refs, {},
                        books_start=START)
    assert je.ready and len(je.entries) == 2, (je.errors, [e.errors for e in je.entries])
    assert je.entries[1].is_adjusting


def test_csv_template_cells_are_formula_safe():
    text = template_bytes("open-items", "csv").decode("utf-8-sig")
    assert not any(cell.startswith(("=", "+", "@")) for line in text.splitlines() for cell in line.split(","))


def test_entry_problem_shared_with_the_je_form():
    body = JournalEntryIn(entry_date=date(2025, 12, 31), lines=[
        {"account_code": "1020", "debit": "1"}, {"account_code": "3010", "credit": "1"}])
    assert "books start" in entry_problem(CHART, body, START)
    ok = body.model_copy(update={"entry_date": START})
    assert entry_problem(CHART, ok, START) == ""
