"""Report exports (XLSX / PDF tables) and the JSON encoding of report data."""

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.services.accounting.books.chart import Chart, standard_chart
from app.services.accounting.books_pdf import Header
from app.services.accounting.compute.books import Entry, Line
from app.services.accounting.compute.reports import columns_for, profit_and_loss, with_change
from app.services.accounting.report_export import engine_error_rows, jsonable, report_pdf, report_xlsx, table_for

D = Decimal
CHART = Chart(standard_chart("corporation"))
HEADER = Header("Bayanihan Hardware Trading Corp.", "123-456-789-00000", "Marikina")
ENTRIES = [Entry("doc:a:sale", date(2026, 10, 5), "SJ", "0006", "=HYPERLINK(\"x\")", "", "Sale", "document",
                 (Line("1100", D("11200")), Line("4010", credit=D("10000")), Line("2100", credit=D("1200"))))]


def test_grid_table_formats_pct_and_indents():
    cols = columns_for(date(2026, 10, 1), date(2026, 12, 31), "total", 12)
    prior = columns_for(date(2026, 7, 1), date(2026, 9, 30), "total", 12)
    report = with_change(profit_and_loss(ENTRIES + [Entry("doc:b:sale", date(2026, 8, 1), "SJ", "1", "", "", "",
                                                          "document", (Line("1100", D("5000")),
                                                                       Line("4010", credit=D("5000"))))],
                                         CHART, cols + prior), 0, 1)
    heads, rows = table_for("pnl", report)
    assert heads[0] == "Account" and heads[-2:] == ["₱ change", "% change"]
    net = next(r for r in rows if r["cells"][0] == "Net sales")
    assert net["cells"][1:] == [D("10000.00"), D("5000.00"), D("5000.00"), "100.00%"] and net["bold"]
    sales = next(r for r in rows if r["cells"][0].strip() == "4010 Sales")
    assert sales["cells"][0].startswith("  ")


def test_xlsx_cells_are_formula_safe():
    rows = [{"cells": ["=cmd|' /C calc'!A0", D("1.50"), date(2026, 10, 5)], "bold": False}]
    data = report_xlsx(HEADER, "General ledger", "Oct 1 – Dec 31, 2026", ["Party", "Amount", "Date"], rows)
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A1"].value == "Bayanihan Hardware Trading Corp."
    values = [c.value for c in ws[6]]
    assert values[0] == "'=cmd|' /C calc'!A0" and values[1] == 1.5


def test_pdf_renders():
    heads, rows = table_for("transactions", {"rows": [{
        "entry_id": "doc:a:sale", "date": date(2026, 10, 5), "book": "SJ", "source": "document", "reference": "0006",
        "party": "Kalayaan", "particulars": "Sale", "accounts": ["1100", "4010"], "account_names": ["AR", "Sales"],
        "amount": D("11200.00")}]})
    data = report_pdf(HEADER, "Transaction list by date", "Oct 2026", heads, rows)
    assert data.startswith(b"%PDF")


def test_jsonable_decimals_and_dates():
    assert jsonable({"a": D("1.5"), "b": [date(2026, 1, 2), None]}) == {"a": "1.50", "b": ["2026-01-02", None]}
    import uuid
    u = uuid.uuid4()
    assert jsonable({"id": u}) == {"id": str(u)}


def test_engine_error_rows():
    rows = engine_error_rows("Unbalanced journal entry je:abc: debits 10.00 ≠ credits 9.00")
    assert rows[0]["entry_id"] == "je:abc" and rows[0]["problem"] == "engine_error"
    assert rows[0]["ref"] == {"kind": "journal_entry", "id": "abc"}
    assert engine_error_rows("The chart of accounts has no account for 'bank'")[0]["entry_id"] == ""
