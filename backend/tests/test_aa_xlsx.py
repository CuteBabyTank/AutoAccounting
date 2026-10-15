from openpyxl import load_workbook

from app.services.accounting.render_xlsx import write_workbook

HEADER = {"company": "Juan's Cloth Trading", "tin": "123-456-789-00000", "period": "2026 · Q3 (Jul–Sep)",
          "generated": "2026-10-05 21:00"}


def _snap(total="120.00", extra_row=False):
    rows = [{"key": "555", "tin": "555", "name": "A", "address": "", "exempt": "0.00", "zero_rated": "0.00",
             "taxable": "1000.00", "gross_sales": "1000.00", "output_tax": "120.00"}]
    if extra_row:
        rows.append({**rows[0], "key": "777", "tin": "777", "name": "B"})
    return {"company": {}, "forms": {"2550Q": {"31A": "1000.00", "26": total}, "1702Q": {"25": "5000"}},
            "ledgers": {"SLS": rows}}


def _write(tmp_path, name="out.xlsx", **kw):
    path = tmp_path / name
    write_workbook(path, header=HEADER, snapshot=kw.pop("snapshot", _snap()), warnings=kw.pop("warnings", []),
                   review_open=kw.pop("review_open", 0), due={"2550Q": "2026-10-26"}, version=kw.pop("version", 1), **kw)
    return load_workbook(path)


def _row(ws, key):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == key)


def _yellow(cell) -> bool:
    return cell.fill is not None and cell.fill.fill_type == "solid" and str(cell.fill.fgColor.rgb).endswith("FFF4B3")


def test_tabs_values_and_number_formats(tmp_path):
    wb = _write(tmp_path)
    assert wb.sheetnames == ["Summary", "2550Q", "1702Q", "SLS"]
    ws = wb["2550Q"]
    assert ws.cell(4, 1).value == "31A"  # spec order: Part IV first
    r = _row(ws, "31A")
    assert ws.cell(r, 3).value == 1000 and ws.cell(r, 3).number_format == "#,##0.00"
    assert ws.cell(_row(ws, "26"), 4).value == "21 + 25"
    assert wb["1702Q"].cell(_row(wb["1702Q"], "25"), 3).number_format == "#,##0"
    assert wb["SLS"].cell(2, 6).value == 1000
    assert not any(_yellow(c) for row in ws.iter_rows() for c in row)


def test_changed_cell_is_yellow_with_comment(tmp_path):
    diff = {"forms": {"2550Q": {"26": ["100.00", "120.00"]}},
            "ledgers": {"SLS": {"added": ["777"], "changed": {}, "removed": ["999"]}}}
    wb = _write(tmp_path, snapshot=_snap(extra_row=True), version=2, diff=diff, parent_version=1)
    ws = wb["2550Q"]
    cell = ws.cell(_row(ws, "26"), 3)
    assert _yellow(cell) and "v1: 100.00 → v2: 120.00" in cell.comment.text
    assert not _yellow(ws.cell(_row(ws, "31A"), 3))
    sls = wb["SLS"]
    new_row = next(r for r in range(2, sls.max_row + 1) if sls.cell(r, 1).value == "777")
    assert all(_yellow(sls.cell(new_row, c)) for c in range(1, 9))
    log = wb["Change log"]
    texts = [str(c.value) for row in log.iter_rows() for c in row if c.value is not None]
    assert "(row removed)" in texts and "(new row)" in texts


def test_clean_copy_has_no_highlights(tmp_path):
    diff = {"forms": {"2550Q": {"26": ["100.00", "120.00"]}}, "ledgers": {}}
    wb = _write(tmp_path, version=2, diff=diff, parent_version=1, clean=True)
    assert "Change log" not in wb.sheetnames
    ws = wb["2550Q"]
    assert not _yellow(ws.cell(_row(ws, "26"), 3)) and ws.cell(_row(ws, "26"), 3).comment is None


def test_summary_has_payables_draft_banner_and_warnings(tmp_path):
    wb = _write(tmp_path, review_open=2, warnings=["Ending inventory was not provided."])
    texts = [str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value is not None]
    assert "DRAFT — 2 document(s) still need review" in texts
    assert "• Ending inventory was not provided." in texts
    assert "2026-10-26" in texts


def test_extracted_text_cannot_become_a_formula(tmp_path):
    evil = '=HYPERLINK("http://x","click")'
    snap = _snap()
    snap["ledgers"]["SLS"][0]["name"] = evil
    diff = {"forms": {}, "ledgers": {"SLS": {"added": [], "changed": {"555": {"name": ["A", evil]}}, "removed": []}}}
    header = {**HEADER, "company": "+cmd"}
    path = tmp_path / "x.xlsx"
    write_workbook(path, header=header, snapshot=snap, warnings=[], review_open=0, due={}, version=2, diff=diff,
                   parent_version=1)
    wb = load_workbook(path)
    cells = [c for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any(c.data_type == "f" for c in cells)
    assert wb["SLS"].cell(2, 2).value == "'" + evil
