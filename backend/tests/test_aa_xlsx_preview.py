"""Workbook preview: displayed text, highlights, bold, and the size caps."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.accounting import xlsx_preview
from app.services.accounting.xlsx_preview import workbook_preview


def _book() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "2550Q"
    ws.append(["Item", "Value"])
    ws["A1"].font = Font(bold=True)
    ws.append(["31A VATable sales", 117390])
    ws["B2"].number_format = "#,##0.00"
    ws["B2"].fill = PatternFill("solid", fgColor="FFF4B3")
    ws.append([None, None])
    wb.create_sheet("QAP").append(["Payee", "Tax"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_cells_keep_their_display_text_highlight_and_bold():
    sheets = workbook_preview(_book())
    assert [s["name"] for s in sheets] == ["2550Q", "QAP"]
    first = sheets[0]
    assert first["rows"][0][0] == {"v": "Item", "bold": True}
    assert first["rows"][1][1] == {"v": "117,390.00", "num": True, "fill": "#FFF4B3"}
    assert len(first["rows"]) == 2  # trailing empty row dropped
    assert first["truncated"] is False


def test_large_sheets_are_capped(monkeypatch):
    monkeypatch.setattr(xlsx_preview, "MAX_ROWS", 1)
    sheets = workbook_preview(_book())
    assert len(sheets[0]["rows"]) == 1 and sheets[0]["truncated"] is True
