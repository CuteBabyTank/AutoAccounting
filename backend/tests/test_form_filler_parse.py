"""Excel parsing tests for the form filler service (compact two-perspective layout)."""
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.services.form_filler_service import (
    PAYEE_KIND, PAYOR_KIND, ParseResult, SHEET_LABELS, build_template, parse_workbook,
)

PAYEE_IDENTITY = ["payee_tin", "payee_name", "payee_registered_address",
                  "payee_zip_code", "payee_foreign_address", "from_date", "to_date"]
PAYOR_IDENTITY = ["payor_tin", "payor_name", "payor_registered_address",
                  "payor_zip_code", "from_date", "to_date"]
ENTITY_TAIL = ["block", "description", "atc", "m1", "m2", "m3", "total", "tax_withheld"]


def _build(kind: str, identity_headers, identity_data, entity_headers,
           entity_rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_LABELS[kind]
    ws.cell(row=1, column=1, value="Payee Form" if kind == PAYEE_KIND else "Payor Form")
    ws.append([])  # placeholder so identity header lands on row 2
    for c, h in enumerate(identity_headers, start=1):
        ws.cell(row=2, column=c, value=h)
    for c, v in enumerate(identity_data, start=1):
        ws.cell(row=3, column=c, value=v)
    ws.append([])  # row 4 blank
    # row 5 left blank (decorative month labels in real template)
    ws.append([])
    for c, h in enumerate(entity_headers, start=1):
        ws.cell(row=6, column=c, value=h)
    r = 7
    for row in entity_rows:
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
        r += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _payee_workbook(entity_rows) -> bytes:
    return _build(
        PAYEE_KIND, PAYEE_IDENTITY,
        ["111-123-323-000", "Fraser", "Quezon City", "1000", "", "01/01/2026", "03/31/2026"],
        ["payor_name", "payor_registered_address", "payor_zip_code", "payor_tin", *ENTITY_TAIL],
        entity_rows,
    )


def _payor_workbook(entity_rows) -> bytes:
    return _build(
        PAYOR_KIND, PAYOR_IDENTITY,
        ["222-222-222-000", "Acme Corp", "Makati", "1226", "01/01/2026", "03/31/2026"],
        ["payee_name", "payee_registered_address", "payee_zip_code", "payee_tin", *ENTITY_TAIL],
        entity_rows,
    )


def test_parse_payee_single_payor() -> None:
    xlsx = _payee_workbook([
        ["Josh Giddy", "Pasig", "1000", "111-111-111-000",
         "ewt", "Professional Fee", "WI010", 3400, 3400, 3400, 10200, 1155],
    ])
    result = parse_workbook(xlsx)
    assert isinstance(result, ParseResult)
    assert result.kind == PAYEE_KIND
    assert result.payee is not None and result.payee.name == "Fraser"
    assert result.payee.tin == "111-123-323-000"
    assert result.period_from == "01/01/2026"
    assert len(result.forms) == 1
    f = result.forms[0]
    assert f.payor.name == "Josh Giddy"
    assert f.payor.tin == "111-111-111-000"
    assert f.payee.name == "Fraser"  # constant identity propagated
    assert len(f.ewt) == 1 and f.ewt[0].atc == "WI010"


def test_blank_name_is_continuation_line_item() -> None:
    xlsx = _payee_workbook([
        ["Josh Giddy", "Pasig", "1000", "111-111-111-000",
         "ewt", "Professional Fee", "WI010", 3400, 3400, 3400, 10200, 1155],
        ["", "", "", "",
         "ewt", "Very Professional Fee", "WI010", 5000, 4000, 2500, 11500, 1150],
        ["Mallow James", "Quezon City", "1000", "222-222-222-000",
         "ewt", "Professional Fee", "WI010", 2000, 2000, 2000, 6000, 930],
    ])
    result = parse_workbook(xlsx)
    assert len(result.forms) == 2  # only two distinct payors
    josh = result.forms[0]
    assert josh.payor.name == "Josh Giddy"
    assert len(josh.ewt) == 2  # continuation attached
    assert josh.ewt[1].description == "Very Professional Fee"
    assert result.forms[1].payor.name == "Mallow James"
    assert len(result.forms[1].ewt) == 1


def test_block_mp_routes_to_money_payments() -> None:
    xlsx = _payee_workbook([
        ["Stephen Yuan", "Quezon City", "1000", "333-333-333-000",
         "mp", "Professional Fee", "WC158", 3300, 3300, 3300, 9900, 660],
    ])
    result = parse_workbook(xlsx)
    f = result.forms[0]
    assert len(f.money_payments) == 1
    assert not f.ewt


def test_parse_payor_form_lists_payees() -> None:
    xlsx = _payor_workbook([
        ["Payee One", "Cebu", "6000", "555-555-555-000",
         "ewt", "Services", "WI010", 1000, 0, 0, 1000, 50],
        ["Payee Two", "Davao", "8000", "666-666-666-000",
         "ewt", "Services", "WI010", 2000, 0, 0, 2000, 100],
    ])
    result = parse_workbook(xlsx)
    assert result.kind == PAYOR_KIND
    assert result.payor is not None and result.payor.name == "Acme Corp"
    assert len(result.forms) == 2
    assert result.forms[0].payee.name == "Payee One"
    assert result.forms[0].payor.name == "Acme Corp"  # constant identity propagated
    assert result.forms[1].payee.tin == "666-666-666-000"


def test_single_period_propagates_to_all_forms() -> None:
    xlsx = _payee_workbook([
        ["A Corp", "Pasig", "1000", "111-111-111-000",
         "ewt", "X", "WI010", 100, 0, 0, 100, 5],
        ["B Corp", "Pasig", "1000", "222-222-222-000",
         "ewt", "Y", "WI010", 200, 0, 0, 200, 10],
    ])
    result = parse_workbook(xlsx)
    assert all(f.period_from == "01/01/2026" and f.period_to == "03/31/2026"
               for f in result.forms)


def test_human_labels_round_trip() -> None:
    """Pretty hand-made headers (JANUARY/QTR/Payor Name) parse via aliases."""
    xlsx = _build(
        PAYEE_KIND, PAYEE_IDENTITY,
        ["111-123-323-000", "Fraser", "QC", "1000", "", "01/01/2026", "03/31/2026"],
        ["Payor Name", "Payor registered address", "payor zip code", "payor tin",
         "block", "description", "atc", "JANUARY", "FEBRUARY", "MARCH", "QTR", "WITHHELD"],
        [["Josh Giddy", "Pasig", "1000", "111-111-111-000",
          "ewt", "Professional Fee", "WI010", 3400, 3400, 3400, 10200, 1155]],
    )
    result = parse_workbook(xlsx)
    assert result.forms[0].payor.name == "Josh Giddy"
    assert result.forms[0].ewt[0].m1 == 3400


def test_detect_falls_back_to_entity_header_without_title() -> None:
    xlsx = _build(
        PAYEE_KIND, PAYEE_IDENTITY,
        ["111-123-323-000", "Fraser", "QC", "1000", "", "01/01/2026", "03/31/2026"],
        ["payor_name", "payor_registered_address", "payor_zip_code", "payor_tin", *ENTITY_TAIL],
        [["Josh", "Pasig", "1000", "111-111-111-000",
          "ewt", "Fee", "WI010", 100, 0, 0, 100, 5]],
    )
    # Wipe the title cell so detection must use the entity header.
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(xlsx))
    wb.active.cell(row=1, column=1, value=None)
    buf = BytesIO()
    wb.save(buf)
    result = parse_workbook(buf.getvalue())
    assert result.kind == PAYEE_KIND


def test_undetectable_workbook_raises() -> None:
    wb = Workbook()
    wb.active.title = "SomethingElse"
    wb.active["A1"] = "Hello"
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        parse_workbook(buf.getvalue())


def test_line_item_before_first_entity_raises() -> None:
    xlsx = _payee_workbook([
        ["", "", "", "", "ewt", "Orphan", "WI010", 100, 0, 0, 100, 5],
    ])
    with pytest.raises(ValueError) as exc:
        parse_workbook(xlsx)
    assert "no payor" in str(exc.value).lower()


def test_entity_excel_row_is_tracked() -> None:
    xlsx = _payee_workbook([
        ["Josh", "Pasig", "1000", "111-111-111-000",
         "ewt", "Fee", "WI010", 100, 0, 0, 100, 5],
    ])
    result = parse_workbook(xlsx)
    # Entity rows start at row 7 in this layout.
    assert result.forms[0].excel_row == 7
    assert result.forms[0].ewt[0].excel_row == 7


@pytest.mark.parametrize("kind", [PAYEE_KIND, PAYOR_KIND])
def test_generated_template_round_trips(kind: str) -> None:
    result = parse_workbook(build_template(kind))
    assert result.kind == kind
    assert len(result.forms) >= 1
    # The sample template demonstrates the continuation rule on the first entity.
    assert len(result.forms[0].ewt) == 2
